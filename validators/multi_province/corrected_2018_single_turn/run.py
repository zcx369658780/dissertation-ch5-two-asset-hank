"""Prepare and execute the one authorized corrected-2018 ordered turn.

``prepare`` is static and creates the immutable runtime payload. ``run`` enters
science exactly once and refuses a second launch in the same evidence root.
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import platform
import sys
import time
from dataclasses import asdict, fields, is_dataclass
from hashlib import sha256
from pathlib import Path
from types import MappingProxyType
from typing import Any, Mapping

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import brentq

REPO = Path(__file__).resolve().parents[3]
for item in (REPO / "src", REPO):
    if str(item) not in sys.path:
        sys.path.insert(0, str(item))

from exports import matlab_faithful_two_asset_ha as oracle
from ch5_two_asset_hank.multi_province import one_turn as one_turn_module
from ch5_two_asset_hank.multi_province.annual import _normalize_province, _xlsx_sheet_rows
from ch5_two_asset_hank.multi_province.one_turn import OneTurnInputs, PreFrozenHouseholdOutputBatch
from ch5_two_asset_hank.multi_province.province_contracts import PROVINCE_ORDER
from ch5_two_asset_hank.multi_province.steady_state import _adapt, _diagnostics, _post_turn_states

CANONICAL_SHA256 = "AEA5A12B5E6474056C1C3EF84BF0156BA88442EF54B0A4FB9C4C6F33CA963F67"
DISTANCE_SHA256 = "26E44D174A8EFFBDCA526D95DA38F0E5883E0C78FDFD036D2DFF1D1FBA5A3566"
CONTRACT = "CH5_ANNUAL_TEMPORAL_CONTRACT_V2_ROLLING10Y_SAMEYEAR_ZT"
VERDICT_PASS = "CORRECTED_2018_SINGLE_TURN_PASS__FULL_ORDERED_TURN_COMPLETED"
VERDICT_HOUSEHOLD = "CORRECTED_2018_SINGLE_TURN_FAIL__HOUSEHOLD_OR_NUMERICAL_BLOCKER"
VERDICT_UPSTREAM = "CORRECTED_2018_SINGLE_TURN_FAIL__UPSTREAM_FIRM_OR_STATE_BLOCKER"
THREAD_ENV = ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS", "NUMEXPR_NUM_THREADS")
LABOR_BRENTQ = brentq


def file_sha256(path: Path) -> str:
    digest = sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def jsonable(value: Any) -> Any:
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value.resolve())
    if is_dataclass(value):
        return {field.name: jsonable(getattr(value, field.name)) for field in fields(value)}
    if isinstance(value, Mapping):
        return {str(key): jsonable(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [jsonable(item) for item in value]
    return value


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(jsonable(payload), stream, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)
        stream.write("\n")


def _rows(workbook, sheet: str) -> list[dict[str, Any]]:
    values = list(workbook[sheet].iter_rows(values_only=True))
    headers = [str(item) for item in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(item is not None for item in row)]


def _column(row: Mapping[str, Any], prefix: str) -> Any:
    matches = [value for key, value in row.items() if key.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"expected one workbook column beginning with {prefix!r}")
    return matches[0]


def accepted_scalars() -> dict[str, float]:
    return {
        "gdp_multiplier": 1000.0, "pop_multiplier": 100.0, "calibration_delta": 0.096,
        "zt_ratio": 1.0, "gov_inv_ratio": 1.0, "max_sigmau": 0.5,
        "rb_gap": 0.07, "rah": 0.09, "ra": 0.09, "nominal_rate": 0.02,
        "rb": 0.02, "rk": 0.1, "wjt": 0.6, "composite_wage": 20.0,
        "transfer_income": 0.1, "inflation": 0.02, "wage_tax": 0.05,
        "initial_at": 2.0, "initial_bt": 1.0, "initial_mt": 0.9,
        "initial_ct": 4.0, "corporate_tax": 0.25,
    }


def source_labor_root(*, wage: float, temp: float, alphac: float = 1.0,
                      alphal: float = 1.0, frisch_l: float = 0.2,
                      ga: float = 2.0, max_search: int = 128) -> tuple[float, tuple[float, float]]:
    values = (wage, temp, alphac, alphal, frisch_l, ga)
    if not all(np.isfinite(value) for value in values):
        raise ValueError("source labor parameters must be finite")
    if wage <= 0 or alphac <= 0 or alphal <= 0 or frisch_l <= 0 or ga <= 0:
        raise ValueError("source labor parameters are outside the frozen positive regime")
    power = ga * frisch_l
    coefficient = (alphac / alphal * wage) ** frisch_l
    x0 = wage ** (frisch_l * (1.0 - ga) / (1.0 + ga * frisch_l))
    boundary = -temp / wage

    def residual(labor: float) -> float:
        base = labor * wage + temp
        if not np.isfinite(labor) or not base > 0:
            raise ValueError("labor residual evaluation outside the real-valued source domain")
        return labor - coefficient * base ** (-power)

    if not x0 > boundary:
        raise ValueError("source x0 is outside the real-valued labor domain")
    at_x0 = residual(x0)
    if at_x0 == 0:
        return x0, (x0, x0)
    if at_x0 < 0:
        lo, hi = x0, max(1.0, 2.0 * x0)
        for _ in range(max_search):
            if residual(hi) > 0:
                break
            hi *= 2.0
        else:
            raise RuntimeError("finite upward labor bracket search exhausted")
    else:
        hi = x0
        lo = 0.5 * (boundary + hi)
        for _ in range(max_search):
            if residual(lo) < 0:
                break
            lo = 0.5 * (boundary + lo)
        else:
            raise RuntimeError("finite interior labor bracket search exhausted")
    root = LABOR_BRENTQ(residual, lo, hi, xtol=1e-14, rtol=1e-14)
    return float(root), (float(lo), float(hi))


def source_initial_arrays(state: Mapping[str, object], grid: Any, params: Any,
                          on_root=None, on_root_return=None) -> tuple[np.ndarray, np.ndarray]:
    shape = (grid.b.size, grid.a.size, grid.z.size)
    labor = np.empty(shape)
    value = np.empty(shape)
    for k, z in enumerate(grid.z):
        for j, a in enumerate(grid.a):
            effective = float(oracle.matlab_faithful_illiquid_return(a, grid.a[-1], float(state["rah"])))
            for i, b in enumerate(grid.b):
                rb = float(state["rb"]) + (float(state["rb_gap"]) if b < 0 else 0.0)
                temp = effective * effective + rb * b + float(state["Tt"])
                wage = (1 - float(state["tau"])) * float(state["w"]) * z
                if on_root is not None:
                    on_root()
                solved_labor, _ = source_labor_root(
                    wage=wage, temp=temp, frisch_l=1.0 / params.phi, ga=params.gamma_c)
                if on_root_return is not None:
                    on_root_return()
                consumption = wage * solved_labor + rb * b + float(state["Tt"])
                labor[i, j, k] = solved_labor
                value[i, j, k] = (
                    consumption ** (1 - params.gamma_c) / (1 - params.gamma_c)
                    - solved_labor ** 6 / 6
                ) / params.rho
    return value, labor


def build_runtime_payload(canonical_workbook: Path, distance_workbook: Path) -> dict[str, Any]:
    canonical_workbook = Path(canonical_workbook)
    distance_workbook = Path(distance_workbook)
    canonical_sha = file_sha256(canonical_workbook)
    if canonical_sha != CANONICAL_SHA256:
        raise ValueError(f"canonical workbook SHA mismatch: {canonical_sha}")
    distance_sha = file_sha256(distance_workbook)
    if distance_sha != DISTANCE_SHA256:
        raise ValueError(f"distance workbook SHA mismatch: {distance_sha}")

    workbook = load_workbook(canonical_workbook, read_only=True, data_only=True)
    try:
        province_rows = _rows(workbook, "PROVINCE_ORDER")
        order = tuple(_normalize_province(row["Province"]) for row in province_rows)
        if order != PROVINCE_ORDER:
            raise ValueError("canonical workbook province order mismatch")
        if [int(row["Python index (0-based)"]) for row in province_rows] != list(range(31)):
            raise ValueError("canonical workbook Python indices mismatch")

        binding_rows = [row for row in _rows(workbook, "ANNUAL_BINDING_2009_2023") if int(row["Steady year"]) == 2018]
        if len(binding_rows) != 1:
            raise ValueError("canonical workbook lacks one 2018 temporal binding")
        binding = binding_rows[0]
        expected_binding = {
            "Steady year": 2018, "Analysis index": 10, "data_MAT index": 10,
            "Level row": 19, "Level year": 2018, "PLM vintage": 19,
            "PLM window start": 2009, "PLM window end": 2018,
            "Window length": 10, "Window type": "ROLLING_10_YEAR",
            "Zt year": 2018, "Contract version": CONTRACT,
        }
        if any(binding[key] != value for key, value in expected_binding.items()):
            raise ValueError("canonical workbook 2018 temporal contract mismatch")

        def annual(sheet: str) -> list[dict[str, Any]]:
            rows = [row for row in _rows(workbook, sheet) if int(next(iter(row.values()))) == 2018]
            rows.sort(key=lambda row: int(row["Python index"]))
            if len(rows) != 31 or tuple(_normalize_province(row["Province"]) for row in rows) != order:
                raise ValueError(f"{sheet} does not contain the ordered 2018 province panel")
            return rows

        gdp_rows = annual("GDP_2000_2023")
        pop_rows = annual("POP_2000_2023")
        cap_rows = annual("PIM_CAPITAL_2000_2023")
        zt_rows = annual("ZT_SAMEYEAR_2009_2023")
        alpha_rows = [row for row in _rows(workbook, "PLM_ALPHA_2009_2023") if int(row["Year"]) == 2018]
        if len(alpha_rows) != 1:
            raise ValueError("canonical workbook lacks one 2018 PLM record")
        alpha = float(alpha_rows[0]["Alpha"])
        gdp_raw = np.array([float(_column(row, "Final-use GDP")) for row in gdp_rows])
        pop_raw = np.array([float(_column(row, "Final-use population")) for row in pop_rows])
        cap_raw = np.array([float(_column(row, "Derived capital")) for row in cap_rows])
        gdp = gdp_raw * 1000.0
        pop = pop_raw * 100.0
        cap = cap_raw * 1000.0
        ind_alpha = np.full(31, alpha)
        ind_zt = np.array([float(_column(row, "IND_Zt static")) for row in zt_rows])
        rebuilt_zt = gdp * cap ** (-ind_alpha) * pop ** (ind_alpha - 1.0)
        if not np.array_equal(ind_zt, rebuilt_zt):
            differences = np.flatnonzero(ind_zt != rebuilt_zt)
            raise ValueError(f"same-year Zt formula identity mismatch at indices {differences.tolist()}")
        if not (gdp_raw[11] == 34010.9 and gdp[11] == 34010900.0 and pop_raw[11] == 6076.0
                and pop[11] == 607600.0 and cap_raw[11] == 1357314108.2013683
                and cap[11] == 1357314108201.3684 and alpha == 0.772866243094144
                and ind_zt[11] == 0.0006934644495858679):
            raise ValueError("accepted Anhui 2018 identity mismatch")
        pcap = cap / pop
        inter_ratio = 0.3 * (pcap - np.min(pcap)) / (np.max(pcap) - np.min(pcap))
    finally:
        workbook.close()

    distance_rows = _xlsx_sheet_rows(distance_workbook, "geom")
    row_axis = tuple(_normalize_province(distance_rows[index][1]) for index in range(2, 33))
    col_axis = tuple(_normalize_province(distance_rows[1][index]) for index in range(2, 33))
    if row_axis != order or col_axis != order:
        raise ValueError("distance workbook axes mismatch")
    distance = np.array([[distance_rows[r][c] for c in range(2, 33)] for r in range(2, 33)], dtype=float)
    sigmau = distance / np.max(distance) * 0.5
    scalars = accepted_scalars()
    states = []
    for index, name in enumerate(order):
        states.append({
            "name": name, "N": pop[index], "alpha": ind_alpha[index], "Zt": ind_zt[index],
            "Kt0": cap[index], "Kt": cap[index], "Kt_prev": cap[index],
            "Lt": pop[index], "Lt_prev": pop[index], "Yt0": gdp[index], "Yt": gdp[index],
            "Zt_1": ind_zt[index], "GovInv": cap[index], "inter_prv_ratio": inter_ratio[index],
            "rb_gap": scalars["rb_gap"], "rah": scalars["rah"], "ra": scalars["ra"],
            "it": scalars["nominal_rate"], "rb": scalars["rb"], "rk": scalars["rk"],
            "wjt": scalars["wjt"], "w": scalars["composite_wage"], "Tt": scalars["transfer_income"],
            "pit": scalars["inflation"], "pit_1": scalars["inflation"], "totalpit": scalars["inflation"],
            "epsilon_pi": 0.0, "tau": scalars["wage_tax"], "At": scalars["initial_at"],
            "Bt": scalars["initial_bt"], "mt": scalars["initial_mt"], "Ct": scalars["initial_ct"],
            "AtTax": 0.0, "GovSurplus": 0.0, "corptau": scalars["corporate_tax"],
            "ramin": 0.02, "ramax": 0.09, "wjtmin": 0.8, "wjtmax": 1.3,
        })
    return {
        "schema": "CH5_CORRECTED_2018_SINGLE_TURN_RUNTIME_INPUT_V1",
        "canonical_workbook": {"sha256": canonical_sha, "bytes": canonical_workbook.stat().st_size},
        "distance_workbook": {"sha256": distance_sha, "bytes": distance_workbook.stat().st_size},
        "temporal_contract": expected_binding, "province_order": list(order),
        "source_rows": {"gdp": [row["Source cell"] for row in gdp_rows],
                        "pop": [row["Source cell"] for row in pop_rows]},
        "vectors": {"gdp_raw": gdp_raw, "gdp": gdp, "pop_raw": pop_raw, "pop": pop,
                    "cap_raw": cap_raw, "cap": cap, "alpha": ind_alpha, "ind_zt": ind_zt,
                    "gov_inv": cap, "inter_province_asset_ratio": inter_ratio},
        "matrices": {"sigmau_destination_origin": sigmau},
        "scalars": scalars, "states": states,
    }


def prepare(canonical_workbook: Path, distance_workbook: Path, evidence_root: Path) -> None:
    root = Path(evidence_root)
    root.mkdir(parents=True, exist_ok=False)
    payload = build_runtime_payload(canonical_workbook, distance_workbook)
    payload_path = root / "runtime_input_payload.json"
    write_json(payload_path, payload)
    write_json(root / "canonical_identity_receipt.json", {
        "schema": "CH5_CORRECTED_2018_CANONICAL_IDENTITY_RECEIPT_V1",
        "status": "PASS", "expected_sha256": CANONICAL_SHA256,
        "actual_sha256": payload["canonical_workbook"]["sha256"],
        "bytes": payload["canonical_workbook"]["bytes"], "scientific_calls": 0,
    })
    write_json(root / "runtime_input_receipt.json", {
        "schema": "CH5_CORRECTED_2018_RUNTIME_INPUT_RECEIPT_V1",
        "runtime_payload_sha256": file_sha256(payload_path),
        "province_count": 31, "province_order": payload["province_order"],
        "contract": CONTRACT, "grid": {"I": 20, "b": [-2.0, 5.0], "J": 20,
                  "a": [0.0, 10.0], "Nz": 2, "z": [0.8, 1.3]},
        "anhui": {key: payload["vectors"][key][11] for key in
                  ("gdp_raw", "gdp", "pop_raw", "pop", "cap_raw", "cap", "alpha", "ind_zt", "gov_inv")},
        "binding_status": "EXISTING_ECONOMIC_MAPPING_PRESERVED", "scientific_calls": 0,
    })
    write_json(root / "source_code_identity.json", {
        "schema": "CH5_CORRECTED_2018_SOURCE_IDENTITY_V1",
        "files": {str(path.relative_to(REPO)).replace("\\", "/"): file_sha256(path) for path in (
            Path(__file__), REPO / "exports/matlab_faithful_two_asset_ha.py",
            REPO / "validators/multi_province/mp4b_python_empirical.py",
            REPO / "src/ch5_two_asset_hank/multi_province/one_turn.py",
            REPO / "src/ch5_two_asset_hank/multi_province/steady_state.py",
            REPO / "src/ch5_two_asset_hank/multi_province/firm.py")},
        "scientific_calls": 0,
    })


def _save_hjb(path: Path, result: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(path)
    np.savez_compressed(path, value=result.value, initial_value=result.initial_value,
        consumption=result.consumption, labor=result.labor, transfer=result.transfer,
        adjustment_cost=result.adjustment_cost, effective_illiquid_return=result.effective_illiquid_return,
        mu_a=result.mu_a, mu_b=result.mu_b, utility=result.utility,
        liquid_label=result.liquid_label, transfer_label=result.transfer_label,
        operator=result.operator.full, post_convergence_operator=result.post_convergence_operator.full,
        iterations=np.int64(result.iterations), converged=np.int8(result.converged),
        convergence_statistic=np.float64(result.convergence_statistic))


def _save_kfe(path: Path, result: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError(path)
    np.savez_compressed(path, original_operator=result.original_operator, transpose=result.transpose,
        contaminated_row_index=np.int64(result.contaminated_row_index), contaminated_matrix=result.contaminated_matrix,
        rhs=result.rhs, raw_solve_vector=result.raw_solve_vector,
        normalization_factor=np.float64(result.normalization_factor), density_vector=result.density_vector,
        density=result.density, db=np.float64(result.db), da=np.float64(result.da),
        cell_weight=np.float64(result.cell_weight), raw_residual_inf=np.float64(result.raw_residual_inf))


def execute(evidence_root: Path) -> int:
    global LABOR_BRENTQ
    root = Path(evidence_root)
    payload_path = root / "runtime_input_payload.json"
    if not payload_path.is_file() or not (root / "runtime_input_receipt.json").is_file():
        raise ValueError("prepared runtime input is missing")
    if (root / "science_started.json").exists():
        raise RuntimeError("scientific execution already started; retry prohibited")
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    if payload["canonical_workbook"]["sha256"] != CANONICAL_SHA256 or payload["temporal_contract"]["Contract version"] != CONTRACT:
        raise ValueError("prepared runtime identity mismatch")
    counters = {
        "scientific_processes": 1, "one_turn_executions": 0, "one_turn_returns": 0,
        "province_updates_attempted": 0, "province_updates_completed": 0,
        "native_initializations_attempted": 0, "native_initializations_returned": 0,
        "labor_roots_attempted": 0, "labor_roots_returned": 0,
        "brentq_calls_attempted": 0, "brentq_calls_returned": 0,
        "household_calls_attempted": 0, "household_calls_returned": 0, "household_calls_failed": 0,
        "hjb_calls": 0, "hjb_returns": 0, "hjb_direct_solves": 0,
        "kfe_calls": 0, "kfe_returns": 0, "kfe_direct_solves": 0,
        "aggregate_calls": 0, "aggregate_returns": 0,
        "migration_calls": 0, "capital_allocation_calls": 0,
        "firm_calls": 0, "firm_returns": 0, "wage_batch_calls": 0,
        "wage_province_outputs": 0, "controller_calls": 0, "adaptation_blocks": 0,
        "scientific_retries": 0,
    }
    write_json(root / "science_started.json", {
        "schema": "CH5_CORRECTED_2018_SCIENCE_LAUNCH_V1", "started_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "runtime_payload_sha256": file_sha256(payload_path), "process_id": os.getpid(),
        "python": sys.version, "platform": platform.platform(), "thread_environment": {key: os.environ.get(key) for key in THREAD_ENV},
        "authorized_one_turn_executions": 1, "scientific_retries": 0,
    })
    started = time.monotonic()
    states = tuple(dict(item) for item in payload["states"])
    grid = oracle.MatlabFaithfulHJBGrid(np.linspace(-2, 5, 20), np.linspace(0, 10, 20),
        np.array([0.8, 1.3]), np.array([[-1/3, 1/3], [1/3, -1/3]]))
    params = oracle.EconomicParams(0.05, 2.0, 5.0, 0.1, 2.0, 1e-6, 0.0, 0.0)
    numerics = oracle.MatlabFaithfulHJBNumerics(1000.0, 1e-7, 100, 1e-12)
    model_params = MappingProxyType({"ga": 2.0, "phi_l": 5.0, "alphal": 1.0, "epsilon": 10.0,
        "theta": 100.0, "delta": 0.025, "istar": 0.015, "rho_pi": 1.25,
        "totalpit": 0.02, "epsilon_pi": 0.0})
    productivity = np.array([float(state["Yt"]) / float(state["Lt"]) for state in states])
    phi = 1.0 + 0.3 * (productivity[:, None] - productivity[None, :]) / (productivity[:, None] + productivity[None, :])
    wedges = np.array(payload["matrices"]["sigmau_destination_origin"], dtype=float)
    outputs: list[dict[str, Any]] = []
    current = {"stage": "PRE_HOUSEHOLD", "province_index": None, "province": None}
    original_spsolve = oracle.linalg.spsolve
    original_brentq = LABOR_BRENTQ

    def counted_spsolve(*args, **kwargs):
        key = "hjb_direct_solves" if current["stage"] == "HJB" else "kfe_direct_solves"
        counters[key] += 1
        return original_spsolve(*args, **kwargs)

    def counted_brentq(*args, **kwargs):
        counters["brentq_calls_attempted"] += 1
        result = original_brentq(*args, **kwargs)
        counters["brentq_calls_returned"] += 1
        return result

    oracle.linalg.spsolve = counted_spsolve
    LABOR_BRENTQ = counted_brentq
    verdict = VERDICT_HOUSEHOLD
    error: dict[str, Any] | None = None
    try:
        for index, state in enumerate(states):
            province = payload["province_order"][index]
            current.update(stage="NATIVE_INITIALIZATION", province_index=index, province=province)
            counters["household_calls_attempted"] += 1
            counters["native_initializations_attempted"] += 1
            def enter_root() -> None:
                counters["labor_roots_attempted"] += 1
            def return_root() -> None:
                counters["labor_roots_returned"] += 1
            initial, labor = source_initial_arrays(state, grid, params, enter_root, return_root)
            counters["native_initializations_returned"] += 1
            current["stage"] = "HJB"
            counters["hjb_calls"] += 1
            hjb = oracle.solve_matlab_faithful_hjb(grid, params, oracle.HouseholdInputs(
                float(state["rah"]), float(state["rb"]), float(state["tau"]),
                np.array([state["w"]]), np.array([0.0]), np.array([1.0])),
                initial, labor, float(state["Tt"]), float(state["rb_gap"]), numerics)
            counters["hjb_returns"] += 1
            province_root = root / "household" / f"p{index:02d}_{province}"
            _save_hjb(province_root / "hjb_return.npz", hjb)
            write_json(province_root / "hjb_receipt.json", {"province_index": index, "province": province,
                "converged": hjb.converged, "iterations": hjb.iterations,
                "convergence_statistic": hjb.convergence_statistic,
                "hjb_direct_solves_cumulative": counters["hjb_direct_solves"],
                "persisted_before_kfe": True})
            current["stage"] = "KFE"
            counters["kfe_calls"] += 1
            kfe = oracle.solve_matlab_faithful_stationary_kfe(hjb.post_convergence_operator.full,
                shape=(20, 20, 2), db=float(grid.b[1] - grid.b[0]), da=float(grid.a[1] - grid.a[0]))
            counters["kfe_returns"] += 1
            _save_kfe(province_root / "kfe_return.npz", kfe)
            current["stage"] = "AGGREGATE"
            counters["aggregate_calls"] += 1
            aggregate = oracle.aggregate_stationary_household(grid, hjb.consumption, hjb.labor, kfe.density)
            counters["aggregate_returns"] += 1
            effective = oracle.matlab_faithful_illiquid_return(grid.a, grid.a[-1], float(state["rah"]))
            at_tax = aggregate.a_ss * float(state["rah"]) - float(
                np.sum(grid.a[None, :, None] * effective[None, :, None] * kfe.density) * kfe.cell_weight)
            result_row = {"province_index": index, "province": province, "hjb_converged": hjb.converged,
                "hjb_iterations": hjb.iterations, "hjb_statistic": hjb.convergence_statistic,
                "kfe_returned": True, "kfe_raw_residual_inf": kfe.raw_residual_inf,
                "C": aggregate.c_ss, "L": aggregate.l_ss, "A": aggregate.a_ss, "B": aggregate.b_ss,
                "A_plus_B": aggregate.a_ss + aggregate.b_ss, "AtTax": at_tax}
            write_json(province_root / "household_return.json", result_row)
            outputs.append(result_row)
            counters["household_calls_returned"] += 1

        batch = PreFrozenHouseholdOutputBatch(
            ct=[row["C"] for row in outputs], household_lt=[row["L"] for row in outputs],
            at=[row["A"] for row in outputs], bt=[row["B"] for row in outputs],
            at_tax=[row["AtTax"] for row in outputs], converged=tuple(row["hjb_converged"] for row in outputs),
            diagnostics=tuple({"hjb_converged": row["hjb_converged"], "hjb_iterations": row["hjb_iterations"],
                "hjb_statistic": row["hjb_statistic"], "kfe_returned": True} for row in outputs))

        original_firm = one_turn_module.evaluate_firm
        original_wage = one_turn_module.composite_household_wages
        original_migration = one_turn_module.reconstruct_migration_labor
        original_capital = one_turn_module.allocate_productive_capital
        def observed_firm(province, *args, **kwargs):
            counters["firm_calls"] += 1
            current.update(stage="FIRM", province=province["name"], province_index=payload["province_order"].index(province["name"]))
            result = original_firm(province, *args, **kwargs)
            counters["firm_returns"] += 1
            return result
        def observed_wage(*args, **kwargs):
            counters["wage_batch_calls"] += 1
            result = original_wage(*args, **kwargs)
            counters["wage_province_outputs"] += len(result)
            return result
        def observed_migration(*args, **kwargs):
            counters["migration_calls"] += 1
            return original_migration(*args, **kwargs)
        def observed_capital(*args, **kwargs):
            counters["capital_allocation_calls"] += 1
            return original_capital(*args, **kwargs)
        one_turn_module.evaluate_firm = observed_firm
        one_turn_module.composite_household_wages = observed_wage
        one_turn_module.reconstruct_migration_labor = observed_migration
        one_turn_module.allocate_productive_capital = observed_capital
        try:
            current.update(stage="ONE_TURN", province=None, province_index=None)
            counters["one_turn_executions"] += 1
            turn = one_turn_module.run_source_faithful_one_turn(OneTurnInputs(
                tuple(payload["province_order"]), states, model_params, phi, wedges, batch))
            counters["one_turn_returns"] += 1
        finally:
            one_turn_module.evaluate_firm = original_firm
            one_turn_module.composite_household_wages = original_wage
            one_turn_module.reconstruct_migration_labor = original_migration
            one_turn_module.allocate_productive_capital = original_capital

        current["stage"] = "CONTROLLER"
        before = _post_turn_states(states, batch, turn)
        diagnostic = _diagnostics(before, batch, np.full(31, 3.0), 1e-9)
        nk_gap, yt_gap, household_count, ra_upper, ra_lower, wage_upper, wage_lower, converged = diagnostic
        counters["controller_calls"] += 1
        adapted, actions = _adapt(before, float(np.max(nk_gap)), True)
        counters["adaptation_blocks"] += 1
        rows = []
        for index, state in enumerate(states):
            firm = turn.firms[index]
            action = actions[index]
            row = {
                "province_index": index, "province": state["name"], "GDP": state["Yt0"], "POP": state["N"],
                "CAP": state["Kt0"], "alpha": state["alpha"], "same_year_Zt": state["Zt"], "GovInv": state["GovInv"],
                "entering_firm_ra": state["ra"], "household_rah": state["rah"], "household_rb": state["rb"],
                "household_composite_wage": state["w"], "entering_Lt_prev": state["Lt_prev"],
                "household_Lt_return": outputs[index]["L"], "destination_lt_supply": turn.migration.lt_supply[index],
                "firm_ra0": firm.ra0, "firm_ra_used": firm.ra, "firm_rk": firm.rk,
                "firm_wage_raw": firm.wt0, "firm_wage_used": firm.wjt,
                "hjb_converged": outputs[index]["hjb_converged"], "hjb_iterations": outputs[index]["hjb_iterations"],
                "hjb_statistic": outputs[index]["hjb_statistic"], "kfe_returned": True,
                "C": outputs[index]["C"], "L": outputs[index]["L"], "A": outputs[index]["A"],
                "B": outputs[index]["B"], "A_plus_B": outputs[index]["A_plus_B"],
                "nk_gap": nk_gap[index], "yt_gap": yt_gap[index], "ra_clipped_lower": firm.ra == state["ramin"] and firm.ra0 < state["ramin"],
                "ra_clipped_upper": firm.ra == state["ramax"] and firm.ra0 > state["ramax"],
                "wage_clipped_lower": firm.wjt == state["wjtmin"] and firm.wt0 < state["wjtmin"],
                "wage_clipped_upper": firm.wjt == state["wjtmax"] and firm.wt0 > state["wjtmax"],
                "zt_adjusted": action.zt_adjusted, "zt_after": action.zt_after,
                "govinv_action": action.govinv_action, "govinv_after": action.govinv_after,
            }
            rows.append(row)
        counters["province_updates_attempted"] = 31
        counters["province_updates_completed"] = 31
        write_json(root / "per_province_observables.json", rows)
        with (root / "per_province_observables.csv").open("x", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
        write_json(root / "controller_observables.json", {
            "household_converged_count": household_count, "all_household_hjb_converged": household_count == 31,
            "max_nk_gap": float(np.max(nk_gap)), "max_nk_gap_province": rows[int(np.argmax(nk_gap))]["province"],
            "max_yt_gap": float(np.max(yt_gap)), "ra_upper_count": ra_upper, "ra_lower_count": ra_lower,
            "wage_upper_count": wage_upper, "wage_lower_count": wage_lower, "source_converged": converged,
            "adaptation_allowed": float(np.max(nk_gap)) < 0.1, "actions": actions,
        })
        write_json(root / "anhui_forensic.json", rows[11])
        verdict = VERDICT_PASS
    except Exception as exc:
        counters["household_calls_failed"] = counters["household_calls_attempted"] - counters["household_calls_returned"]
        error = {"type": type(exc).__name__, "message": str(exc), "stage": current["stage"],
                 "province_index": current["province_index"], "province": current["province"]}
        if current["stage"] in ("ONE_TURN", "FIRM", "CONTROLLER"):
            verdict = VERDICT_UPSTREAM
        write_json(root / "scientific_failure.json", error)
    finally:
        oracle.linalg.spsolve = original_spsolve
        LABOR_BRENTQ = original_brentq
    counters["household_calls_failed"] = counters["household_calls_attempted"] - counters["household_calls_returned"]
    write_json(root / "call_ledger.json", {"schema": "CH5_CORRECTED_2018_SINGLE_TURN_CALL_LEDGER_V1",
        "counts": counters, "failed_calls_are_counted": True, "elapsed_seconds": time.monotonic() - started,
        "scientific_retry_count": 0, "matlab_calls": 0, "second_turns": 0, "steady_state_loops": 0,
        "ge_calls": 0, "annual_model_calls": 0, "irf_calls": 0, "results_calls": 0})
    write_json(root / "terminal_result.json", {"schema": "CH5_CORRECTED_2018_SINGLE_TURN_TERMINAL_V1",
        "verdict": verdict, "error": error, "full_ordered_turn_completed": verdict == VERDICT_PASS,
        "downstream_steady_state_authorized": "NO", "results_eligible": False})
    return 0 if verdict == VERDICT_PASS else 2


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    prepare_parser = sub.add_parser("prepare")
    prepare_parser.add_argument("canonical_workbook", type=Path)
    prepare_parser.add_argument("distance_workbook", type=Path)
    prepare_parser.add_argument("evidence_root", type=Path)
    run_parser = sub.add_parser("run")
    run_parser.add_argument("evidence_root", type=Path)
    args = parser.parse_args(argv)
    if args.command == "prepare":
        prepare(args.canonical_workbook, args.distance_workbook, args.evidence_root)
        return 0
    return execute(args.evidence_root)


if __name__ == "__main__":
    raise SystemExit(main())
