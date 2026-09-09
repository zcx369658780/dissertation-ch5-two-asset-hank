"""Build the bounded official-data closure package without scientific calls."""

from __future__ import annotations

import csv
import json
import shutil
import sys
from datetime import date
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from core import NOT_REQUIRED_INVESTMENT_YEAR, REQUIRED_INVESTMENT_YEARS, relative_difference, sha256, transformed_value


SOURCE_ROOT = Path(r"D:\MatlabProgram\2023年12月2日 多省份神经网络HANK")
RAW = SOURCE_ROOT / "2000年后各省数据.xlsx"
FILLED = SOURCE_ROOT / "2000年后各省数据_填充NA.xlsx"
EXPECTED_HASHES = {
    RAW.name: "09814A45D933B2685A35238A15C0C7BB501F00A63597796B3CADCE15C230ECB3",
    FILLED.name: "C826B01B6C124EAAADC063DFC2D5510E50E72ED85BB34848F28AB318E4B88929",
}
RETRIEVAL_DATE = date(2026, 9, 9).isoformat()
AH2018_URL = "https://www.ah.gov.cn/zfsj/tjgblmdz/sjtjgb/554177011.html"
AH2019_URL = "https://www.ah.gov.cn/zfsj/tjgblmdz/sjtjgb/554177001.html"
NBS2019_BASE = "https://www.stats.gov.cn/sj/ndsj/2019/"


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8", newline="\n")


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"no rows for {path}")
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def fetch(url: str, destination: Path) -> dict:
    request = Request(url, headers={"User-Agent": "Mozilla/5.0 official-data-audit/1.0"})
    last_error = None
    for attempt in range(1, 4):
        try:
            with urlopen(request, timeout=30) as response:
                data = response.read()
                status = getattr(response, "status", 200)
                content_type = response.headers.get("Content-Type")
            destination.write_bytes(data)
            return {"status": status, "content_type": content_type, "bytes": len(data), "sha256": sha256(destination), "error": None, "web_attempts": attempt}
        except (HTTPError, URLError, TimeoutError) as error:
            last_error = error
            if isinstance(error, HTTPError):
                break
    return {"status": getattr(last_error, "code", None), "content_type": None, "bytes": 0, "sha256": None, "error": f"{type(last_error).__name__}: {last_error}", "web_attempts": attempt}


def compact_text(path: Path) -> str:
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    return " ".join(soup.get_text(" ", strip=True).split())


def excerpt(text: str, needle: str, before: int = 180, after: int = 520) -> str:
    offset = text.find(needle)
    if offset < 0:
        raise AssertionError(f"required official text not found: {needle}")
    return text[max(0, offset - before) : offset + after]


def workbook_rows() -> tuple[list[dict], dict]:
    raw = load_workbook(RAW, read_only=True, data_only=True)
    filled = load_workbook(FILLED, read_only=True, data_only=True)
    try:
        assert raw["GDP"]["N1"].value == "安徽省"
        assert filled["GDP"]["N1"].value == "安徽省"
        values = {}
        for variable, sheet, cell in (("GDP", "GDP", "N20"), ("POP", "常住人口", "N20"), ("CAP", "总资本存量", "N20")):
            book = filled if variable == "CAP" else raw
            values[variable] = float(book[sheet][cell].value)
        rows = []
        for year in range(2000, 2019):
            row = year - 1998
            raw_value = raw["固定资产投资额"].cell(row, 14).value
            filled_value = filled["固定资产投资额"].cell(row, 14).value
            required = year in REQUIRED_INVESTMENT_YEARS
            if year == 2011:
                comparability = "STATISTICAL_DEFINITION_BREAK"
                caveat = "Official bulletin records a 2011 scope change to projects with planned total investment >= RMB 5 million plus real-estate development investment."
            elif year == NOT_REQUIRED_INVESTMENT_YEAR:
                comparability = "NOT_REQUIRED_BY_RECURRENCE"
                caveat = "K2018 ends with I2017 under the frozen recurrence; I2018 is retained only for ledger completeness."
            else:
                comparability = "SOURCE_NOT_OBTAINED"
                caveat = "No same-definition official absolute observation was automatically obtained."
            rows.append({
                "year": year,
                "workbook_value": raw_value,
                "workbook_status": "ORIGINAL_OBSERVED_UNCHANGED" if raw_value == filled_value else "WORKBOOK_VALUE_CHANGED",
                "official_value": "",
                "official_unit": "",
                "source_id": "AH2018_BULLETIN_SCOPE_NOTE" if year == 2011 else "",
                "comparability_status": comparability,
                "conversion": "none; workbook unit is 万元",
                "accepted_for_candidate_chain": False,
                "required_by_K2018": required,
                "caveat": caveat,
            })
        return rows, values
    finally:
        raw.close()
        filled.close()


def main(evidence_root: str, repo_report_dir: str) -> None:
    evidence = Path(evidence_root)
    repo = Path(repo_report_dir)
    evidence.mkdir(parents=True, exist_ok=False)
    repo.mkdir(parents=True, exist_ok=True)
    downloads = evidence / "official_downloads"
    downloads.mkdir()

    source_specs = [
        ("AH2018_BULLETIN", AH2018_URL, downloads / "anhui_2018_statistical_bulletin.html"),
        ("AH2019_BULLETIN_REVISION_NOTE", AH2019_URL, downloads / "anhui_2019_statistical_bulletin.html"),
        ("NBS2019_YEARBOOK_INDEX", NBS2019_BASE + "indexch.htm", downloads / "nbs_yearbook_2019_index.html"),
        ("NBS2019_YEARBOOK_LEFT", NBS2019_BASE + "left.htm", downloads / "nbs_yearbook_2019_left.html"),
        ("NBS2019_TABLE_2_6_IMAGE", NBS2019_BASE + "html/C0206.jpg", downloads / "nbs_2019_C0206.jpg"),
        ("NBS2019_TABLE_3_9_IMAGE", NBS2019_BASE + "html/C0309.jpg", downloads / "nbs_2019_C0309.jpg"),
        ("NBS2019_TABLE_10_1_IMAGE", NBS2019_BASE + "html/C1001.jpg", downloads / "nbs_2019_C1001.jpg"),
        ("NBS2019_FIXED_INVESTMENT_NOTES", NBS2019_BASE + "html/sm10.htm", downloads / "nbs_2019_sm10.htm"),
    ]
    inventory = []
    for source_id, url, destination in source_specs:
        result = fetch(url, destination)
        inventory.append({"source_id": source_id, "institution": "安徽省统计局、国家统计局安徽调查总队" if source_id.startswith("AH") else "国家统计局 / 中国统计出版社", "publication_database": "安徽省国民经济和社会发展统计公报" if source_id.startswith("AH") else "中国统计年鉴2019", "url": url, "retrieval_date": RETRIEVAL_DATE, **result})

    by_id = {row["source_id"]: row for row in inventory}
    for required in ("AH2018_BULLETIN", "AH2019_BULLETIN_REVISION_NOTE"):
        if by_id[required]["status"] != 200:
            raise RuntimeError(f"required structured official source unavailable: {required}")
    text18 = compact_text(downloads / "anhui_2018_statistical_bulletin.html")
    text19 = compact_text(downloads / "anhui_2019_statistical_bulletin.html")
    excerpts = {
        "source_id": "AH2018_BULLETIN_SCOPE_NOTE",
        "gdp_initial_publication": excerpt(text18, "30006.82"),
        "resident_population_initial_publication": excerpt(text18, "6323.6"),
        "investment_growth_only": excerpt(text18, "11.8%"),
        "investment_scope_break": excerpt(text18, "计划总投资 500 万元"),
        "gdp_revision_notice": excerpt(text19, "第四次全国经济普查"),
    }
    write_json(evidence / "official_text_excerpts.json", excerpts)
    write_csv(repo / "official_source_inventory.csv", inventory)
    source_hash = {row["source_id"]: row["sha256"] for row in inventory}
    observations = [
        {"source_id": "AH2018_BULLETIN", "source_institution": "安徽省统计局、国家统计局安徽调查总队", "publication_database": "安徽省2018年国民经济和社会发展统计公报", "table_title_indicator": "2018年全省生产总值及增速 / 地区生产总值", "province": "安徽省", "year": 2018, "raw_value": 30006.82, "unit": "亿元", "current_constant_price_wording": "absolute value at current prices; growth at comparable prices", "retrieval_date": RETRIEVAL_DATE, "stable_url_publication_id": AH2018_URL, "download_sha256": source_hash["AH2018_BULLETIN"], "page_table_number": "HTML section 一、综合", "revision_conflict_notes": "Preliminary statistics; later 2019 bulletin announces historical GDP revisions.", "use_status": "OFFICIAL_OBSERVATION_OBTAINED__NOT_FINAL_IDENTITY"},
        {"source_id": "AH2018_BULLETIN", "source_institution": "安徽省统计局、国家统计局安徽调查总队", "publication_database": "安徽省2018年国民经济和社会发展统计公报", "table_title_indicator": "2018年末全省常住人口及构成 / 年末常住人口", "province": "安徽省", "year": 2018, "raw_value": 6323.6, "unit": "万人", "current_constant_price_wording": "not applicable", "retrieval_date": RETRIEVAL_DATE, "stable_url_publication_id": AH2018_URL, "download_sha256": source_hash["AH2018_BULLETIN"], "page_table_number": "HTML section 一、综合", "revision_conflict_notes": "Preliminary statistics; differs from provisional workbook 6076 and later official revision lineage was not obtained.", "use_status": "OFFICIAL_OBSERVATION_OBTAINED__NOT_FINAL_IDENTITY"},
        {"source_id": "AH2018_BULLETIN_SCOPE_NOTE", "source_institution": "安徽省统计局、国家统计局安徽调查总队", "publication_database": "安徽省2018年国民经济和社会发展统计公报", "table_title_indicator": "固定资产投资统计范围脚注", "province": "安徽省", "year": 2011, "raw_value": "", "unit": "计划总投资500万元及以上项目和房地产开发投资", "current_constant_price_wording": "scope/coverage definition", "retrieval_date": RETRIEVAL_DATE, "stable_url_publication_id": AH2018_URL, "download_sha256": source_hash["AH2018_BULLETIN"], "page_table_number": "HTML note [6][7]", "revision_conflict_notes": "2011 national statistical-system reform creates a definition break in the required 2000-2017 chain.", "use_status": "OFFICIAL_METHOD_NOTE_ACCEPTED"},
        {"source_id": "AH2019_BULLETIN_REVISION_NOTE", "source_institution": "安徽省统计局、国家统计局安徽调查总队", "publication_database": "安徽省2019年国民经济和社会发展统计公报", "table_title_indicator": "地区生产总值历史数据修订脚注", "province": "安徽省", "year": "historical series including 2018", "raw_value": "", "unit": "亿元", "current_constant_price_wording": "absolute GDP at current prices; historical series revised after fourth national economic census", "retrieval_date": RETRIEVAL_DATE, "stable_url_publication_id": AH2019_URL, "download_sha256": source_hash["AH2019_BULLETIN_REVISION_NOTE"], "page_table_number": "HTML note [2]", "revision_conflict_notes": "The page states that historical GDP was revised but does not expose the exact revised 2018 value.", "use_status": "OFFICIAL_REVISION_NOTICE_ACCEPTED__VALUE_MISSING"},
        {"source_id": "NBS2019_TABLE_IMAGES", "source_institution": "国家统计局 / 中国统计出版社", "publication_database": "中国统计年鉴2019", "table_title_indicator": "2-6 分地区年末人口数; 3-9 地区生产总值(2018年); 10-1 全社会固定资产投资", "province": "安徽省 / 全国", "year": 2018, "raw_value": "", "unit": "table-specific", "current_constant_price_wording": "not extracted", "retrieval_date": RETRIEVAL_DATE, "stable_url_publication_id": NBS2019_BASE + "indexch.htm", "download_sha256": ";".join(filter(None, (source_hash.get("NBS2019_TABLE_2_6_IMAGE"), source_hash.get("NBS2019_TABLE_3_9_IMAGE"), source_hash.get("NBS2019_TABLE_10_1_IMAGE")))), "page_table_number": "tables 2-6, 3-9, 10-1", "revision_conflict_notes": "Official tables are image-only. No OCR or visual transcription was used; Owner manual extraction is requested.", "use_status": "OFFICIAL_SOURCE_LOCATED__STRUCTURED_VALUE_NOT_OBTAINED"},
    ]
    write_csv(repo / "official_observation_ledger.csv", observations)

    for path in (RAW, FILLED):
        actual = sha256(path)
        if actual != EXPECTED_HASHES[path.name]:
            raise RuntimeError(f"source hash mismatch: {path.name} {actual}")
    chain, workbook = workbook_rows()
    write_csv(repo / "anhui_investment_chain_2000_2018.csv", chain)

    identities = [
        {
            "variable": "GDP",
            "province": "安徽省",
            "calendar_year": 2018,
            "provisional_workbook_value": workbook["GDP"],
            "provisional_unit": "亿元",
            "official_value_obtained": 30006.82,
            "official_unit": "亿元",
            "official_vintage": "2018 statistical bulletin; preliminary statistics; current-price absolute GDP",
            "source_id": "AH2018_BULLETIN",
            "absolute_difference_provisional_minus_obtained": round(workbook["GDP"] - 30006.82, 10),
            "relative_difference_provisional_minus_obtained": relative_difference(workbook["GDP"], 30006.82),
            "identity_status": "OFFICIAL_SOURCE_NOT_OBTAINED",
            "closure_caveat": "The 2019 official bulletin says historical GDP was revised after the fourth national economic census, but the exact revised 2018 official value was not present in the obtained structured pages.",
        },
        {
            "variable": "resident_population",
            "province": "安徽省",
            "calendar_year": 2018,
            "provisional_workbook_value": workbook["POP"],
            "provisional_unit": "万人",
            "official_value_obtained": 6323.6,
            "official_unit": "万人",
            "official_vintage": "2018 statistical bulletin; preliminary statistics; year-end resident population",
            "source_id": "AH2018_BULLETIN",
            "absolute_difference_provisional_minus_obtained": round(workbook["POP"] - 6323.6, 10),
            "relative_difference_provisional_minus_obtained": relative_difference(workbook["POP"], 6323.6),
            "identity_status": "OFFICIAL_SOURCE_NOT_OBTAINED",
            "closure_caveat": "The obtained official structured source differs from the provisional workbook; no later official table establishing the workbook value and revision chronology was automatically obtained.",
        },
    ]
    write_csv(repo / "anhui_2018_gdp_population_identity.csv", identities)

    comparison = {
        "schema_version": "CH5_MP4C_2018_OFFICIAL_IDENTITY_COMPARISON_V1",
        "temporal_contract": "CH5_ANNUAL_TEMPORAL_CONTRACT_V2_ROLLING10Y_SAMEYEAR_ZT",
        "analysis_index": 10,
        "data_mat_index": 10,
        "level_row": 19,
        "steady_year": 2018,
        "plm_vintage": 19,
        "plm_window": {"start": 2009, "end": 2018, "length": 10, "type": "rolling"},
        "same_year_zt_year": 2018,
        "alpha": 0.772866243094144,
        "current_provisional": {
            "GDP_raw_yi_yuan": workbook["GDP"], "POP_raw_wan_people": workbook["POP"], "CAP_workbook_derived": workbook["CAP"],
            "GDP_transformed": transformed_value("GDP", workbook["GDP"]), "POP_transformed": transformed_value("POP", workbook["POP"]), "CAP_transformed": transformed_value("CAP", workbook["CAP"]),
        },
        "official_observations_obtained": {"GDP_preliminary_2018_yi_yuan": 30006.82, "resident_population_preliminary_2018_wan_people": 6323.6},
        "candidate_verified": {"GDP": None, "POP": None, "CAP": None, "Zt": None},
        "candidate_package_warranted": False,
        "candidate_package_status": "NO_CANDIDATE_PACKAGE__OFFICIAL_IDENTITIES_AND_INVESTMENT_CHAIN_INCOMPLETE",
        "capital_label_if_future_chain_closes": "MODEL_DERIVED_FROM_OFFICIAL_INVESTMENT",
    }
    write_json(repo / "candidate_2018_input_comparison.json", comparison)

    requests = [
        {"priority": "P0", "variable": "GDP", "province": "安徽省", "year": "2018", "unit": "亿元, current-price absolute value", "current_provisional": workbook["GDP"], "official_table_or_publication": "Post-fourth-economic-census revised Anhui historical GDP table", "required_fields": "2018 revised GDP, current-price wording, table/page, revision vintage", "reason": "2018 bulletin value 30006.82 is preliminary; 2019 bulletin announces historical revision but obtained structured page lacks exact revised 2018 value"},
        {"priority": "P0", "variable": "年末常住人口", "province": "安徽省", "year": "2018", "unit": "万人", "current_provisional": workbook["POP"], "official_table_or_publication": "Anhui Statistical Yearbook or NBS table with revised historical resident-population series", "required_fields": "2018 year-end resident population, survey/census revision basis, table/page", "reason": "Official 2018 bulletin reports 6323.6, while workbook has 6076; later official identity and revision chronology are missing"},
        {"priority": "P0", "variable": "固定资产投资额", "province": "安徽省", "year": "2000-2017", "unit": "万元 or explicitly convertible official unit", "current_provisional": "see anhui_investment_chain_2000_2018.csv", "official_table_or_publication": "Official Anhui yearbook tables with absolute annual investment and methodological notes", "required_fields": "each annual absolute value 2000-2017, unit, price basis, coverage, 2011 linkage or comparable backcast", "reason": "K2018 requires I2000-I2017; 2018 bulletin documents a 2011 coverage break and publishes only growth for 2018"},
    ]
    write_csv(repo / "manual_official_data_request.csv", requests)

    closure = {
        "primary_verdict": "PARTIAL_OFFICIAL_IDENTITY__MANUAL_DATA_REQUIRED",
        "GDP_identity_status": "OFFICIAL_SOURCE_NOT_OBTAINED",
        "population_identity_status": "OFFICIAL_SOURCE_NOT_OBTAINED",
        "investment_chain_status": "CAPITAL_CHAIN_OFFICIAL_CLOSURE_INCOMPLETE",
        "capital_stock_status": "NOT_COMPUTED__INCOMPLETE_OR_NONCOMPARABLE_OFFICIAL_INVESTMENT_CHAIN",
        "candidate_V2_input_status": "NO_CANDIDATE_PACKAGE__OFFICIAL_IDENTITIES_AND_INVESTMENT_CHAIN_INCOMPLETE",
        "official_source_completeness": "PARTIAL_STRUCTURED_OFFICIAL_EVIDENCE__MANUAL_TABLE_EXTRACTION_REQUIRED",
        "scientific_calls": 0,
        "Results_eligible": False,
        "recurrence_identity": {"K0": "I0 / 0.1", "Kt": "(1 - 0.096) * K(t-1) + I(t-1)", "K2018_required_investment_years": list(REQUIRED_INVESTMENT_YEARS), "I2018": "NOT_REQUIRED_BY_RECURRENCE"},
        "source_hashes": {RAW.name: sha256(RAW), FILLED.name: sha256(FILLED)},
        "downloaded_official_source_count": sum(row["status"] == 200 for row in inventory),
        "download_failures": [row for row in inventory if row["status"] != 200],
    }
    write_json(repo / "closure_status.json", closure)
    write_json(repo / "call_ledger.json", {
        "MATLAB": 0, "household_HJB_KFE": 0, "root_brentq_direct_iterative_eigen": 0,
        "firm_one_turn_controller": 0, "stationary_GE_annual_model": 0, "IRF_Results": 0,
        "allowed_static_web_spreadsheet_hash_arithmetic_operations_only": True,
    })
    shutil.copy2(repo / "closure_status.json", evidence / "closure_status.json")
    shutil.copy2(repo / "anhui_investment_chain_2000_2018.csv", evidence / "anhui_investment_chain_2000_2018.csv")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
