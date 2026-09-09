"""Read-only Excel/MAT/source audit. It contains no model or solve entry point."""
import csv, hashlib, json, math, os, re, sys, zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
import h5py
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from contract import annual_indices, classify_fill, missing_runs, parse_year

ROOT=Path(r"D:\MatlabProgram\2023年12月2日 多省份神经网络HANK")
RAW=ROOT/"2000年后各省数据.xlsx"; FILLED=ROOT/"2000年后各省数据_填充NA.xlsx"; RAW2024=ROOT/"2024年数据原始版.xlsx"; MAT=ROOT/"数据估计结果_1000_100_0.mat"
SOURCES=[RAW,FILLED,RAW2024,MAT,ROOT/"load_GDPdata.m",ROOT/"It_to_Kt.m",ROOT/"mpHANK_equilibrium_2000.m",ROOT/"multi_prov_HANK_12sts.m",ROOT/"R语言估计结果_plm估计.xlsx"]
RAW_SHEETS=["GDP","固定资产投资额","总人口数","常住人口","就业人数","第一产业固定资产投资额","第一产业从业人员数","第一产业占GDP的比重","第二产业固定资产投资额","第二产业从业人员数","第二产业占GDP比重","第三产业固定资产投资额","第三产业从业人员数","第三产业占GDP的比重"]
UNITS={"GDP":"亿元","固定资产投资额":"万元","总人口数":"万人","常住人口":"万人","就业人数":"万人","第一产业固定资产投资额":"亿元","第一产业从业人员数":"万人","第一产业占GDP的比重":"%","第二产业固定资产投资额":"亿元","第二产业从业人员数":"万人","第二产业占GDP比重":"%","第三产业固定资产投资额":"亿元","第三产业从业人员数":"万人","第三产业占GDP的比重":"%"}
WIDE_HEADERS={"GDP":"GDP(亿元)","固定资产投资额":"固定资产投资额(万元)","总人口数":"总人口数(万人)","常住人口":"常住人口(万人)","就业人数":"就业人数(万人)","第一产业固定资产投资额":"第一产业固定资产投资额(亿元)","第一产业从业人员数":"第一产业从业人员数(万人)","第一产业占GDP的比重":"第一产业占GDP的比重(%)","第二产业固定资产投资额":"第二产业固定资产投资额(亿元)","第二产业从业人员数":"第二产业从业人员数(万人)","第二产业占GDP比重":"第二产业占GDP比重(%)","第三产业固定资产投资额":"第三产业固定资产投资额(亿元)","第三产业从业人员数":"第三产业从业人员数(万人)","第三产业占GDP的比重":"第三产业占GDP的比重(%)"}

def sha(path):
 d=hashlib.sha256()
 with Path(path).open("rb") as f:
  for b in iter(lambda:f.read(1024*1024),b""): d.update(b)
 return d.hexdigest().upper()
def write_json(p,v): Path(p).write_text(json.dumps(v,ensure_ascii=False,allow_nan=False,indent=2)+"\n",encoding="utf-8",newline="\n")
def norm_name(x): return str(x or "").removesuffix("省").removesuffix("市")
def numeric(x): return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(float(x))

def ooxml_numeric_cells(path, sheet_name):
    main="http://schemas.openxmlformats.org/spreadsheetml/2006/main"; rel="http://schemas.openxmlformats.org/officeDocument/2006/relationships"; pkg="http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path) as z:
        book=ET.fromstring(z.read("xl/workbook.xml")); target_id=next(x.attrib[f"{{{rel}}}id"] for x in book.find(f"{{{main}}}sheets") if x.attrib["name"]==sheet_name)
        rels=ET.fromstring(z.read("xl/_rels/workbook.xml.rels")); target=next(x.attrib["Target"] for x in rels if x.attrib["Id"]==target_id).lstrip("/")
        if not target.startswith("xl/"): target="xl/"+target
        sheet=ET.fromstring(z.read(target)); result=[]
        for cell in sheet.iter(f"{{{main}}}c"):
            value=cell.find(f"{{{main}}}v")
            if value is not None:
                try: result.append((cell.attrib["r"],float(value.text)))
                except (TypeError,ValueError): pass
        return result

def workbook_matrix(path,sheet):
 wb=load_workbook(path,read_only=False,data_only=True); ws=wb[sheet]; years=[parse_year(ws.cell(r,1).value) for r in range(2,26)]; provinces=[ws.cell(1,c).value for c in range(3,34)]; values=np.empty((24,31),object)
 for r in range(24):
  for c in range(31): values[r,c]=ws.cell(r+2,c+3).value
 wb.close(); return years,provinces,values

def mat_char(f,ref): return "".join(chr(int(x)) for x in np.asarray(f[ref][()]).reshape(-1))
def mat_cell(f,g,field,index):
    value=np.asarray(f[g[field][()][index,0]][()]).T
    if value.dtype.fields and set(value.dtype.fields)=={"real","imag"}:
        value=value["real"]+1j*value["imag"]
    return value

def main(evidence,repo_out):
 evidence=Path(evidence); repo_out=Path(repo_out); evidence.mkdir(parents=True,exist_ok=False); repo_out.mkdir(parents=True,exist_ok=True)
 inventory=[]
 for p in SOURCES:
  st=p.stat(); inventory.append({"path":str(p.resolve()),"bytes":st.st_size,"mtime_ns":st.st_mtime_ns,"sha256":sha(p)})
 write_json(evidence/"source_inventory.json",inventory)
 raw_book=load_workbook(RAW,read_only=False,data_only=True); filled_book=load_workbook(FILLED,read_only=False,data_only=True)
 province_order=[raw_book["GDP"].cell(1,c).value for c in range(3,34)]; years=[parse_year(raw_book["GDP"].cell(r,1).value) for r in range(2,26)]; assert len(province_order)==31 and province_order[11]=="安徽省" and years==list(range(2000,2024))
 ledger_path=evidence/"cell_level_ledger.csv"; fields=["variable","raw_workbook","filled_workbook","sheet","cell","province","calendar_year","raw_value","filled_value","changed","missing_original","classification","transformed_model_value","unit_transformation","confidence_caveat"]
 summaries={}; anomalies=[]
 with ledger_path.open("w",encoding="utf-8-sig",newline="") as stream:
  writer=csv.DictWriter(stream,fieldnames=fields); writer.writeheader()
  for sheet in RAW_SHEETS:
   rw=raw_book[sheet]; fw=filled_book[sheet]; classes=Counter(); miss_by_province=defaultdict(list); zeros=0; jumps=[]
   for c0,province in enumerate(province_order):
    series=[rw.cell(r,c0+3).value for r in range(2,26)]
    for r0,year in enumerate(years):
     raw=series[r0]; filled=fw.cell(r0+2,c0+3).value; before=any(numeric(x) for x in series[:r0]); after=any(numeric(x) for x in series[r0+1:]); cls=classify_fill(raw if numeric(raw) else None,filled if numeric(filled) else None,before,after); classes[cls]+=1
     if not numeric(raw): miss_by_province[province].append(year)
     if numeric(filled) and float(filled)<=0: zeros+=1; anomalies.append({"type":"ZERO_OR_NONPOSITIVE","variable":sheet,"province":province,"year":year,"cell":f"{get_column_letter(c0+3)}{r0+2}","value":filled})
     transformed=None; trans=UNITS[sheet]
     if sheet=="GDP" and numeric(filled): transformed=float(filled)*1000; trans="亿元 * 1000"
     if sheet=="常住人口" and numeric(filled): transformed=float(filled)*100; trans="万人 * 100"
     writer.writerow({"variable":sheet,"raw_workbook":RAW.name,"filled_workbook":FILLED.name,"sheet":sheet,"cell":f"{get_column_letter(c0+3)}{r0+2}","province":province,"calendar_year":year,"raw_value":raw,"filled_value":filled,"changed":not (numeric(raw) and numeric(filled) and float(raw)==float(filled)),"missing_original":not numeric(raw),"classification":cls,"transformed_model_value":transformed,"unit_transformation":trans,"confidence_caveat":"makima is source-declared only for missing/non-numeric raw values; nonmissing changes remain unresolved"})
    vals=np.array([float(fw.cell(r,c0+3).value) if numeric(fw.cell(r,c0+3).value) else np.nan for r in range(2,26)])
    for i in range(1,len(vals)):
     if vals[i]>0 and vals[i-1]>0:
      ratio=vals[i]/vals[i-1]
      if ratio>2 or ratio<.5: jumps.append({"province":province,"year":years[i],"ratio":float(ratio)})
   runs=[]
   for province,ys in miss_by_province.items():
    mask=[y in ys for y in years]
    for a,b,n in missing_runs([None if x else 1 for x in mask]): runs.append({"province":province,"start":years[a],"end":years[b],"length":n,"endpoint":a==0 or b==23})
   summaries[sheet]={"unit":UNITS[sheet],"cells":24*31,"class_counts":dict(classes),"missing_runs":runs,"nonpositive_filled":zeros,"abrupt_ratio_gt2_or_lt0p5":jumps}
 raw_book.close(); filled_book.close(); write_json(evidence/"panel_quality.json",summaries)
 # Verify filled capital stocks against the source recurrence using delta=.096.
 cap_checks={}
 for invest,cap in (("固定资产投资额","总资本存量"),("第一产业固定资产投资额","第一产业资本存量"),("第二产业固定资产投资额","第二产业资本存量"),("第三产业固定资产投资额","第三产业资本存量")):
  _,_,iv=workbook_matrix(FILLED,invest); _,_,kv=workbook_matrix(FILLED,cap); calc=np.zeros((24,31)); arr=np.asarray(iv,dtype=float); calc[0]=arr[0]/.1
  for i in range(1,24): calc[i]=(1-.096)*calc[i-1]+arr[i-1]
  diff=calc-np.asarray(kv,dtype=float); cap_checks[cap]={"max_abs_difference":float(np.max(abs(diff))),"exact_equal_cells":int(np.sum(calc==np.asarray(kv,dtype=float))),"cells":744}
 # Cache entry mydata2{10}, total-industry cell 4.
 with h5py.File(MAT,"r") as f:
  g=f[f["mydata2"][()][9,0]]; cache={k:mat_cell(f,g,k,3) for k in ("GDP","CAP","POP","log_pgdp","log_pcap","IND_alpha","IND_Zt")}; cache_names=[mat_char(f,x) for x in g["prvname"][()].reshape(-1)]; scalars={k:float(np.asarray(g[k][()]).reshape(-1)[0]) for k in ("GDP_multiplier","POP_multiplier","delta")}
 _,_,fgdp=workbook_matrix(FILLED,"GDP"); _,_,fpop=workbook_matrix(FILLED,"常住人口"); _,_,fcap=workbook_matrix(FILLED,"总资本存量"); _,_,rgdp=workbook_matrix(RAW,"GDP"); _,_,rpop=workbook_matrix(RAW,"常住人口")
 expected={"GDP":np.asarray(fgdp,dtype=float)*1000,"CAP":np.asarray(fcap,dtype=float)*1000,"POP":np.asarray(fpop,dtype=float)*100,"log_pgdp":np.log((np.asarray(fgdp,dtype=float)*1000/(np.asarray(fpop,dtype=float)*100)).astype(complex)),"log_pcap":np.log((np.asarray(fcap,dtype=float)*1000/(np.asarray(fpop,dtype=float)*100)).astype(complex))}
 cache_checks={k:{"shape":list(cache[k].shape),"max_abs_difference_vs_filled_transform":float(np.max(abs(cache[k]-v))),"exact_equal_cells":int(np.sum(cache[k]==v)),"nonzero_imaginary_cells":int(np.sum(np.imag(cache[k])!=0))} for k,v in expected.items()}
 # 2018-labelled route actually selects data row 10 = workbook year 2009.
 idx=annual_indices(10); ah=11; row_used=9; row_intended=18
 rbook=load_workbook(RAW,read_only=False,data_only=True); fbook=load_workbook(FILLED,read_only=False,data_only=True)
 line=[]
 def add(field,source_sheet,source_cell,actual_year,raw_value,filled_value,cache_value,transform,status,caveat): line.append({"field":field,"source_sheet":source_sheet,"source_cell":source_cell,"province":"安徽省","claimed_calendar_year":2018,"actual_source_year":actual_year,"raw_value":raw_value,"filled_value":filled_value,"cache_or_final_value":cache_value,"transformation":transform,"status":status,"caveat":caveat})
 for field,sheet,mult in (("Yt0/GDP","GDP",1000),("N/POP","常住人口",100)):
  add(field,sheet,"N11",2009,rbook[sheet]["N11"].value,fbook[sheet]["N11"].value,float(cache[field.split('/')[1] if field.startswith('N/') else 'GDP'][row_used,ah]),f"filled * {mult}","MATERIAL_YEAR_MISALIGNMENT","2018-labelled runtime uses row10/2009; intended 2018 is N20")
 add("Kt0/CAP","总资本存量","N11",2009,None,fbook["总资本存量"]["N11"].value,float(cache["CAP"][row_used,ah]),"makima-filled investment -> It_to_Kt(delta=.096) -> *1000","MATERIAL_YEAR_MISALIGNMENT","derived CAP uses 2000-2008 investment history; 2018 intended value is N20")
 add("pgdp/log_pgdp","GDP÷常住人口","N11",2009,None,None,float(np.real(cache["log_pgdp"][row_used,ah])),"log(GDP*1000/(POP*100))","MATERIAL_YEAR_MISALIGNMENT","derived from 2009 levels; imaginary part is zero for this cell")
 add("pcap/log_pcap","总资本存量÷常住人口","N11",2009,None,None,float(np.real(cache["log_pcap"][row_used,ah])),"log(CAP*1000/(POP*100))","MATERIAL_YEAR_MISALIGNMENT","derived from 2009 levels; imaginary part is zero for this cell")
 # Regression coefficient last numeric cell and Zt uses row21 = 2020.
 sheet="总面板回归系数_19_行业4"; nums=ooxml_numeric_cells(ROOT/"R语言估计结果_plm估计.xlsx",sheet); alpha_cell,alpha=nums[-1]
 add("alpha",sheet,alpha_cell,"regression_vintage_19",None,None,float(cache["IND_alpha"][0,ah]),"last numeric coefficient; copied to all provinces","REGRESSION_DERIVED_SHARED","vintage key19; same alpha for 31 provinces")
 add("Zt",f"GDP/CAP/POP + {sheet}","N22",2020,rbook["GDP"]["N22"].value,fbook["GDP"]["N22"].value,float(cache["IND_Zt"][0,ah]),"2020 transformed levels: GDP*CAP^(-alpha)*POP^(alpha-1)","MIXED_YEAR_CALIBRATION","2018-labelled state combines 2009 levels with technology built from 2020 levels")
 intended={"Yt0/GDP":("GDP!N20",rbook["GDP"]["N20"].value,fbook["GDP"]["N20"].value,float(fbook["GDP"]["N20"].value)*1000),"N/POP":("常住人口!N20",rbook["常住人口"]["N20"].value,fbook["常住人口"]["N20"].value,float(fbook["常住人口"]["N20"].value)*100),"Kt0/CAP":("总资本存量!N20",None,fbook["总资本存量"]["N20"].value,float(fbook["总资本存量"]["N20"].value)*1000),"pgdp/log_pgdp":("GDP!N20÷常住人口!N20",None,None,float(np.log(float(fbook["GDP"]["N20"].value)*1000/(float(fbook["常住人口"]["N20"].value)*100)))),"pcap/log_pcap":("总资本存量!N20÷常住人口!N20",None,None,float(np.log(float(fbook["总资本存量"]["N20"].value)*1000/(float(fbook["常住人口"]["N20"].value)*100))))}
 for row in line:
  cell,raw18,filled18,trans18=intended.get(row["field"],("NOT_APPLICABLE",None,None,None)); row.update(intended_2018_cell=cell,intended_2018_raw=raw18,intended_2018_filled=filled18,intended_2018_transformed=trans18)
 rbook.close(); fbook.close()
 # Confirm published runtime input vectors match cache row10 while its metadata claims workbook row19.
 runtime_path=Path(r"D:\ProjectTemp\ch5-mp4c-full-annual-batch-runtime-cache-20260902-002\year_2018\calendar_2018_matlab_runtime_cache_input.json"); runtime=json.loads(runtime_path.read_text(encoding="utf-8")); vector_match={k:bool(np.array_equal(np.asarray(runtime["vectors"][k]),cache[k][row_used])) for k in ("GDP","CAP","POP","log_pgdp","log_pcap")}; metadata_claim=runtime["binding"]["workbook_data_row_index"]
 complex_pcap=[{"year":2000+int(i),"province":cache_names[int(j)],"CAP":float(cache["CAP"][i,j]),"log_pcap_real":float(np.real(cache["log_pcap"][i,j])),"log_pcap_imag":float(np.imag(cache["log_pcap"][i,j]))} for i,j in zip(*np.where(np.imag(cache["log_pcap"])!=0))]
 cache_audit={"readable":"HDF5_V7P3_READ_ONLY","top_variable":"mydata2","entry":"mydata2{10}","province_names":cache_names,"scalars":scalars,"workbook_transform_checks":cache_checks,"complex_log_pcap_cells":complex_pcap,"capital_recurrence_checks":cap_checks,"annual_index_contract":idx,"runtime_input_path":str(runtime_path),"runtime_input_sha256":sha(runtime_path),"runtime_vector_exact_match_cache_row10":vector_match,"runtime_metadata_workbook_data_row_index":metadata_claim,"metadata_conflict":"metadata says row19/calendar2018, but source path strings, cache vectors, and exact values select row10/calendar2009"}
 write_json(evidence/"cache_audit.json",cache_audit)
 # 2024 workbook cross-check using explicit Sheet0 headers.
 wb24=load_workbook(RAW2024,read_only=False,data_only=True); wbraw24=load_workbook(RAW,read_only=False,data_only=True); ws=wb24["Sheet0"]; wide={};
 for c in range(2,ws.max_column+1): wide[(ws.cell(2,c).value,parse_year(ws.cell(3,c).value))]=c
 compare24={}
 for sheet in RAW_SHEETS:
  key=WIDE_HEADERS[sheet]; matched=changed=missing=0
  rw=wbraw24[sheet]
  for year in years:
   c=wide.get((key,year))
   if not c: continue
   for pi,province in enumerate(province_order):
    a=rw.cell(year-2000+2,pi+3).value; b=ws.cell(pi+5,c).value
    if not numeric(a) or not numeric(b): missing+=1
    else: matched+=1; changed+=int(float(a)!=float(b))
  compare24[sheet]={"schema_mapping":key,"numeric_comparable_cells":matched,"different_cells":changed,"missing_either":missing,"status":"COMPARABLE_BY_EXPLICIT_HEADER" if matched else "NOT_COMPARABLE"}
 wb24.close(); wbraw24.close(); write_json(evidence/"workbook_2024_comparison.json",compare24)
 # Write targeted repo artifacts and escalation list.
 with (repo_out/"anhui_2018_lineage.csv").open("w",encoding="utf-8-sig",newline="") as f: w=csv.DictWriter(f,fieldnames=line[0].keys()); w.writeheader(); w.writerows(line)
 anomaly_rows=[]
 for sheet,x in summaries.items():
  for run in x["missing_runs"]: anomaly_rows.append({"severity":"DATA_QUALITY","variable":sheet,"province":run["province"],"year":f"{run['start']}-{run['end']}","issue":"MISSING_RUN_ENDPOINT" if run["endpoint"] else "MISSING_RUN_INTERIOR","detail":f"length={run['length']}"})
 anomaly_rows += [{"severity":"MATERIAL_DATA_QUALITY","variable":"CAP/log_pcap","province":x["province"],"year":x["year"],"issue":"NEGATIVE_DERIVED_CAPITAL_COMPLEX_LOG","detail":f"CAP={x['CAP']}; log imaginary=pi"} for x in complex_pcap]
 material_lineage=[{"severity":"MATERIAL_LINEAGE","variable":"GDP/CAP/POP/logs","province":"安徽省","year":"2018 label -> 2009 data","issue":"YEAR_INDEX_MISALIGNMENT","detail":"ii=10 makes filename2018 but data_year=10 selects workbook row10/calendar2009"},{"severity":"MATERIAL_LINEAGE","variable":"Zt/alpha","province":"安徽省","year":"2018 label mixes 2009 and 2020","issue":"MIXED_YEAR_CALIBRATION","detail":"regression vintage19 alpha; Zt uses fixed workbook row21/calendar2020"},{"severity":"MATERIAL_LINEAGE","variable":"runtime input metadata","province":"all 31","year":"2018","issue":"METADATA_VALUE_CONTRADICTION","detail":"metadata workbook_data_row_index=19 but vectors and source paths exactly match row10"}]
 anomaly_rows += material_lineage
 with (repo_out/"anomalies.csv").open("w",encoding="utf-8-sig",newline="") as f: w=csv.DictWriter(f,fieldnames=anomaly_rows[0].keys()); w.writeheader(); w.writerows(anomaly_rows)
 requests=[
  {"priority":"P0","variable":"GDP","province":"安徽省","year":2018,"unit":"亿元","current_raw":rgdp[row_intended,ah],"current_filled":fgdp[row_intended,ah],"reason":"runtime labelled2018 consumes 2009 row; verify intended 2018 level","suggested_official_source":"国家统计局分地区年度GDP / 安徽统计年鉴"},
  {"priority":"P0","variable":"常住人口","province":"安徽省","year":2018,"unit":"万人","current_raw":rpop[row_intended,ah],"current_filled":fpop[row_intended,ah],"reason":"runtime labelled2018 consumes 2009 row","suggested_official_source":"国家统计局人口抽样调查 / 安徽统计年鉴"},
  {"priority":"P0","variable":"资本存量/固定资产投资链","province":"安徽省","year":"2000-2018","unit":"投资万元; 资本派生单位","current_raw":"资本存量无原始cell；由固定资产投资额N2:N19递推","current_filled":fcap[row_intended,ah],"reason":"2018 capital is recurrence-derived; runtime instead consumes 2009 stock","suggested_official_source":"国家统计局固定资产投资 / 安徽统计年鉴，并核验资本存量构造"},
  {"priority":"P1","variable":"alpha/Zt","province":"安徽省/31省","year":"vintage19 and 2020 anchor","unit":"share/index","current_raw":alpha,"current_filled":float(cache["IND_Zt"][0,ah]),"reason":"2018-labelled state mixes 2009 levels, common alpha and 2020 Zt anchor","suggested_official_source":"保存R回归设计、系数表与官方GDP/CAP/POP口径说明"},
 ]
 requests.extend({"priority":"P1","variable":"固定资产投资额→总资本存量","province":x["province"],"year":x["year"],"unit":"投资万元; 资本派生单位","current_raw":"原始投资存在缺失/非数值，见cell ledger","current_filled":x["CAP"],"reason":"makima/endpoint-filled investment recurrence yields negative capital and complex log_pcap","suggested_official_source":"国家统计局固定资产投资 / 对应省统计年鉴"} for x in complex_pcap)
 with (repo_out/"official_data_request.csv").open("w",encoding="utf-8-sig",newline="") as f: w=csv.DictWriter(f,fieldnames=requests[0].keys()); w.writeheader(); w.writerows(requests)
 compact={"classification":["MATERIAL_DATA_QUALITY_OR_LINEAGE_ISSUE_FOUND","PARTIAL_EVIDENCE_NEEDS_OFFICIAL_DATA"],"province_order":province_order,"anhui":{"zero_based":11,"one_based":12,"excel_column":"N","header":"安徽省"},"years":years,"data_contract":{"source_raw_sheets":RAW_SHEETS,"actual_annual_fields":["GDP{4}","CAP{4}","POP{4}","log_pgdp{4}","log_pcap{4}","IND_alpha{4}","IND_Zt{4}"],"multipliers":{"GDP":1000,"CAP":1000,"POP":100},"fill":"fillmissing(...,'makima') only when filled workbook absent; existing filled workbook is read directly","capital":"It_to_Kt: K0=I0/.1; Kt=(1-.096)K(t-1)+I(t-1)","smooth_method":0,"reg_method":0},"year_contract":idx,"material_findings":anomaly_rows[-3:],"panel_summary":{k:{"class_counts":v["class_counts"],"missing_run_count":len(v["missing_runs"]),"nonpositive_filled":v["nonpositive_filled"],"jump_count":len(v["abrupt_ratio_gt2_or_lt0p5"])} for k,v in summaries.items()},"cache_audit":cache_audit,"workbook_2024_comparison":compare24,"zero_model_call_ledger":{"MATLAB_model":0,"Python_HJB_KFE_household":0,"firm_controller_GE_annual_IRF_Results":0,"root_direct_iterative_eigen_model_solve":0},"Results_eligible":False}
 write_json(repo_out/"summary.json",compact); write_json(repo_out/"source_inventory.json",inventory); write_json(repo_out/"cache_audit.json",cache_audit); write_json(repo_out/"workbook_2024_comparison.json",compare24); write_json(repo_out/"panel_quality_summary.json",compact["panel_summary"])
 write_json(evidence/"terminal.json",{"completed":True,"classification":compact["classification"],"full_cell_ledger_rows":24*31*len(RAW_SHEETS),"model_calls":compact["zero_model_call_ledger"]}); print(json.dumps({"classification":compact["classification"],"anhui":compact["anhui"],"year_contract":idx,"ledger_rows":24*31*len(RAW_SHEETS)},ensure_ascii=False))
if __name__=="__main__": main(sys.argv[1],sys.argv[2])
