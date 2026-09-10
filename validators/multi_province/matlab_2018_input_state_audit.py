from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Iterable

import h5py
import numpy as np
from openpyxl import load_workbook


EXPECTED_CANONICAL_SHA256 = (
    "AEA5A12B5E6474056C1C3EF84BF0156BA88442EF54B0A4FB9C4C6F33CA963F67"
)
NOT_AVAILABLE = "NOT_AVAILABLE_WITHOUT_SCIENTIFIC_RERUN"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def read_sheet_records(path: Path, sheet: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook[sheet].values)
    finally:
        workbook.close()
    header = [str(value) for value in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:] if any(v is not None for v in row)]


def read_panel_year(path: Path, sheet: str, year: int) -> tuple[list[str], list[float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook[sheet].values)
    finally:
        workbook.close()
    headers = [normalize_province(str(value)) for value in rows[0][2:33]]
    target = next(row for row in rows[1:] if str(row[0]).strip() == f"{year}年")
    return headers, [float(value) for value in target[2:33]]


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    rows = list(rows)
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def normalize_province(value: str) -> str:
    for suffix in ("省", "市", "自治区", "壮族自治区", "回族自治区", "维吾尔自治区"):
        if value.endswith(suffix):
            value = value[: -len(suffix)]
            break
    return value


def decode_matlab_char(dataset: h5py.Dataset) -> str:
    return "".join(chr(int(value)) for value in np.asarray(dataset).reshape(-1))


def matlab_cell_numeric(h5: h5py.File, group: h5py.Group, field: str, matlab_cell_index: int) -> np.ndarray:
    # MATLAB v7.3 stores numeric matrices transposed in HDF5.
    ref = group[field][matlab_cell_index - 1, 0]
    return np.asarray(h5[ref], dtype=float).T


def read_matlab_cache(path: Path, analysis_ii: int) -> dict[str, Any]:
    with h5py.File(path, "r") as h5:
        group = h5[h5["mydata2"][analysis_ii - 1, 0]]
        province_names = [decode_matlab_char(h5[group["prvname"][i, 0]]) for i in range(31)]
        values = {
            "province_names": province_names,
            "GDP": matlab_cell_numeric(h5, group, "GDP", 4),
            "CAP": matlab_cell_numeric(h5, group, "CAP", 4),
            "POP": matlab_cell_numeric(h5, group, "POP", 4),
            "alpha": matlab_cell_numeric(h5, group, "IND_alpha", 4),
            "Zt": matlab_cell_numeric(h5, group, "IND_Zt", 4),
            "GDP_multiplier": float(np.asarray(group["GDP_multiplier"])[0, 0]),
            "POP_multiplier": float(np.asarray(group["POP_multiplier"])[0, 0]),
            "delta": float(np.asarray(group["delta"])[0, 0]),
        }
    return values


def read_csv_by(path: Path, key: str, predicate=lambda row: True) -> dict[str, dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = [row for row in csv.DictReader(fh) if predicate(row)]
    return {normalize_province(row[key]): row for row in rows}


def f(value: Any) -> float:
    return float(value)


def rel_diff(left: float, right: float) -> float:
    return (left - right) / right


def finite_or_text(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return NOT_AVAILABLE
    return value


def make_markdown_ledger(path: Path, ledger: list[dict[str, Any]]) -> None:
    columns = [
        "province_index",
        "province_name",
        "MATLAB_level_year",
        "canonical_year",
        "MATLAB_GDP_raw_亿元",
        "canonical_GDP_raw_亿元",
        "GDP_relative_difference",
        "MATLAB_POP_raw_万人",
        "canonical_POP_raw_万人",
        "MATLAB_alpha",
        "canonical_alpha",
        "MATLAB_Zt_fixed_2020_level",
        "canonical_same_year_Zt",
        "Zt_ratio",
        "MATLAB_raw_capital_万元",
        "canonical_PIM_capital_万元",
        "capital_ratio",
        "initial_ra",
        "initial_rah",
        "initial_rb",
        "initial_wjt",
        "initial_household_wage",
    ]
    lines = [
        "# 31-province human-readable comparison ledger",
        "",
        "All differences compare the original MATLAB 2018-labelled initialization with the corrected canonical 2018 object. Scientific/model calls: 0.",
        "",
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join(["---"] * len(columns)) + " |",
    ]
    for row in ledger:
        vals = []
        for column in columns:
            value = row[column]
            if isinstance(value, float):
                vals.append(format(value, ".12g"))
            else:
                vals.append(str(value))
        lines.append("| " + " | ".join(vals) + " |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def top_rows(ledger: list[dict[str, Any]], metric: str, n: int = 5, absolute: bool = True) -> list[dict[str, Any]]:
    ranked = sorted(
        ledger,
        key=lambda row: abs(f(row[metric])) if absolute else f(row[metric]),
        reverse=True,
    )[:n]
    return [
        {"metric": metric, "rank": rank, "province": row["province_name"], "value": row[metric]}
        for rank, row in enumerate(ranked, start=1)
    ]


def build(args: argparse.Namespace) -> dict[str, Any]:
    repo = Path(args.repo).resolve()
    source = Path(args.matlab_root).resolve()
    canonical = Path(args.canonical).resolve()
    output = repo / args.output
    evidence = Path(args.evidence).resolve()
    output.mkdir(parents=True, exist_ok=True)
    evidence.mkdir(parents=True, exist_ok=False)

    actual_canonical_hash = sha256(canonical)
    if actual_canonical_hash != EXPECTED_CANONICAL_SHA256:
        raise RuntimeError(
            f"canonical hash mismatch: expected {EXPECTED_CANONICAL_SHA256}, got {actual_canonical_hash}"
        )

    processed = source / "2000年后各省数据_填充NA.xlsx"
    raw = source / "2000年后各省数据.xlsx"
    plm = source / "R语言估计结果_plm估计.xlsx"
    cache = source / "数据估计结果_1000_100_0.mat"
    steady_cache = source / "Multi_Province_12sts_2018.mat"
    distance_workbook = source / "中国各省省会地理距离矩阵.xlsx"
    cache_data = read_matlab_cache(cache, analysis_ii=10)

    order = read_sheet_records(canonical, "PROVINCE_ORDER")
    gdp = {
        (int(row["Year"]), row["Province"]): row
        for row in read_sheet_records(canonical, "GDP_2000_2023")
    }
    pop = {
        (int(row["Year"]), row["Province"]): row
        for row in read_sheet_records(canonical, "POP_2000_2023")
    }
    capital = {
        (int(row["Capital year"]), row["Province"]): row
        for row in read_sheet_records(canonical, "PIM_CAPITAL_2000_2023")
    }
    zt = {
        (int(row["Year"]), row["Province"]): row
        for row in read_sheet_records(canonical, "ZT_SAMEYEAR_2009_2023")
    }
    alpha_row = next(
        row for row in read_sheet_records(canonical, "PLM_ALPHA_2009_2023") if int(row["Year"]) == 2018
    )
    binding = next(
        row for row in read_sheet_records(canonical, "ANNUAL_BINDING_2009_2023") if int(row["Steady year"]) == 2018
    )

    corrected = read_csv_by(
        repo / "reports/mp4c_2018_corrected_single_turn_20260909/per_province_observables.csv",
        "province",
    )
    legacy_household = read_csv_by(
        repo / "reports/2018_observable_prefix_replay_20260908/household_entries.csv",
        "province",
        lambda row: row["step"] == "1",
    )
    legacy_firm = read_csv_by(
        repo / "reports/2018_observable_prefix_replay_20260908/firm_prices_and_operands.csv",
        "province",
        lambda row: row["step"] == "1",
    )

    expected_names = [row["Province"] for row in order]
    cache_names = [normalize_province(name) for name in cache_data["province_names"]]
    if cache_names != expected_names:
        raise AssertionError(f"province order mismatch: cache={cache_names}, canonical={expected_names}")
    if cache_data["GDP"].shape != (24, 31) or cache_data["CAP"].shape != (24, 31):
        raise AssertionError("unexpected cached panel shape")

    workbook_gdp_names, workbook_gdp_2009 = read_panel_year(processed, "GDP", 2009)
    workbook_cap_names, workbook_cap_2009 = read_panel_year(processed, "总资本存量", 2009)
    workbook_pop_names, workbook_pop_2009 = read_panel_year(processed, "常住人口", 2009)
    if workbook_gdp_names != expected_names or workbook_cap_names != expected_names or workbook_pop_names != expected_names:
        raise AssertionError("processed workbook province order does not match canonical order")
    cache_workbook_diffs = {
        "GDP_2009_after_x1000": float(np.max(np.abs(cache_data["GDP"][9] - np.asarray(workbook_gdp_2009) * 1000.0))),
        "CAP_2009_after_x1000": float(np.max(np.abs(cache_data["CAP"][9] - np.asarray(workbook_cap_2009) * 1000.0))),
        "POP_2009_after_x100": float(np.max(np.abs(cache_data["POP"][9] - np.asarray(workbook_pop_2009) * 100.0))),
    }
    if any(value != 0.0 for value in cache_workbook_diffs.values()):
        raise AssertionError(f"cache/workbook mismatch: {cache_workbook_diffs}")
    zt_reconstructed = (
        cache_data["GDP"][20]
        * np.power(cache_data["CAP"][20], -cache_data["alpha"][0])
        * np.power(cache_data["POP"][20], cache_data["alpha"][0] - 1.0)
    )
    zt_reconstruction_max_abs_diff = float(np.max(np.abs(zt_reconstructed - cache_data["Zt"][0])))
    if zt_reconstruction_max_abs_diff > 1e-15:
        raise AssertionError(f"cached Zt does not reconstruct from row21 formula: {zt_reconstruction_max_abs_diff}")

    ledger: list[dict[str, Any]] = []
    for zero_index, order_row in enumerate(order):
        province = order_row["Province"]
        matlab_gdp_transformed = f(cache_data["GDP"][9, zero_index])
        matlab_pop_transformed = f(cache_data["POP"][9, zero_index])
        matlab_cap_transformed = f(cache_data["CAP"][9, zero_index])
        matlab_alpha = f(cache_data["alpha"][0, zero_index])
        matlab_zt = f(cache_data["Zt"][0, zero_index])
        canonical_gdp_raw = f(gdp[(2018, province)]["Final-use GDP (亿元)"])
        canonical_pop_raw = f(pop[(2018, province)]["Final-use population (万人)"])
        canonical_cap_raw = f(capital[(2018, province)]["Derived capital (万元)"])
        canonical_alpha = f(alpha_row["Alpha"])
        canonical_zt = f(zt[(2018, province)]["IND_Zt static"])
        matlab_gdp_raw = matlab_gdp_transformed / cache_data["GDP_multiplier"]
        matlab_pop_raw = matlab_pop_transformed / cache_data["POP_multiplier"]
        matlab_cap_raw = matlab_cap_transformed / cache_data["GDP_multiplier"]
        canonical_gdp_transformed = canonical_gdp_raw * cache_data["GDP_multiplier"]
        canonical_pop_transformed = canonical_pop_raw * cache_data["POP_multiplier"]
        canonical_cap_transformed = canonical_cap_raw * cache_data["GDP_multiplier"]

        corrected_row = corrected[province]
        legacy_h = legacy_household[province]
        legacy_f = legacy_firm[province]
        corrected_at = f(corrected_row["A"])
        corrected_bt = f(corrected_row["B"])
        legacy_kt_supply = f(legacy_f["kt_supply"])
        legacy_govinv = f(legacy_f["input_GovInv"])

        row = {
            "province_index": zero_index + 1,
            "province_name": province,
            "source_excel_column": order_row["Source Excel column"],
            "MATLAB_steady_label": 2018,
            "MATLAB_analysis_ii": 10,
            "MATLAB_data_year_index": 10,
            "MATLAB_level_row": 10,
            "MATLAB_level_year": 2009,
            "MATLAB_Zt_source_row": 21,
            "MATLAB_Zt_source_year": 2020,
            "MATLAB_PLM_vintage": 19,
            "canonical_year": 2018,
            "canonical_level_row": int(binding["Level row"]),
            "MATLAB_GDP_raw_亿元": matlab_gdp_raw,
            "MATLAB_GDP_transformed": matlab_gdp_transformed,
            "canonical_GDP_raw_亿元": canonical_gdp_raw,
            "canonical_GDP_transformed": canonical_gdp_transformed,
            "GDP_abs_diff_亿元": matlab_gdp_raw - canonical_gdp_raw,
            "GDP_relative_difference": rel_diff(matlab_gdp_raw, canonical_gdp_raw),
            "GDP_classification": "YEAR_OR_VINTAGE_MISMATCH",
            "MATLAB_POP_raw_万人": matlab_pop_raw,
            "MATLAB_POP_transformed": matlab_pop_transformed,
            "canonical_POP_raw_万人": canonical_pop_raw,
            "canonical_POP_transformed": canonical_pop_transformed,
            "POP_diff_万人": matlab_pop_raw - canonical_pop_raw,
            "POP_relative_difference": rel_diff(matlab_pop_raw, canonical_pop_raw),
            "POP_classification": "YEAR_OR_VINTAGE_MISMATCH",
            "MATLAB_alpha": matlab_alpha,
            "canonical_alpha": canonical_alpha,
            "alpha_diff": matlab_alpha - canonical_alpha,
            "alpha_classification": "EXACT_MATCH" if matlab_alpha == canonical_alpha else "CACHE_OR_PROVENANCE_AMBIGUITY",
            "MATLAB_Zt_fixed_2020_level": matlab_zt,
            "canonical_same_year_Zt": canonical_zt,
            "Zt_ratio": matlab_zt / canonical_zt,
            "Zt_log_ratio": math.log(matlab_zt / canonical_zt),
            "Zt_relative_difference": rel_diff(matlab_zt, canonical_zt),
            "Zt_classification": "YEAR_OR_VINTAGE_MISMATCH",
            "MATLAB_raw_capital_万元": matlab_cap_raw,
            "MATLAB_transformed_capital": matlab_cap_transformed,
            "canonical_PIM_capital_万元": canonical_cap_raw,
            "canonical_transformed_capital": canonical_cap_transformed,
            "capital_ratio": matlab_cap_raw / canonical_cap_raw,
            "capital_relative_difference": rel_diff(matlab_cap_raw, canonical_cap_raw),
            "capital_classification": "YEAR_OR_VINTAGE_MISMATCH",
            "MATLAB_initial_GovInv": matlab_cap_transformed,
            "canonical_initial_GovInv": canonical_cap_transformed,
            "GovInv_ratio_MATLAB_to_canonical": matlab_cap_transformed / canonical_cap_transformed,
            "GovInv_to_initial_Kt0": 1.0,
            "initial_source_At": 2.0,
            "initial_source_Bt": 1.0,
            "initial_source_At_times_N": 2.0 * matlab_pop_transformed,
            "corrected_turn1_persisted_At": corrected_at,
            "corrected_turn1_persisted_Bt": corrected_bt,
            "corrected_turn1_At_times_N": corrected_at * canonical_pop_transformed,
            "corrected_turn1_persisted_Kt_supply": NOT_AVAILABLE,
            "corrected_turn1_Kt_total": NOT_AVAILABLE,
            "legacy_replay_turn1_persisted_At": f(legacy_h["a_ss"]),
            "legacy_replay_turn1_persisted_Bt": f(legacy_h["b_ss"]),
            "legacy_replay_turn1_persisted_Kt_supply": legacy_kt_supply,
            "legacy_replay_turn1_GovInv": legacy_govinv,
            "legacy_replay_turn1_Kt_total": legacy_kt_supply + legacy_govinv,
            "legacy_replay_turn1_GovInv_to_private_supply": legacy_govinv / legacy_kt_supply,
            "initial_ra": 0.09,
            "initial_rah": 0.09,
            "initial_rb": 0.02,
            "initial_wjt": 0.6,
            "initial_household_wage": 20.0,
            "ra_bounds": "[0.02,0.09]",
            "wjt_bounds": "[0.8,1.3]",
            "initial_ra_at_upper_bound": True,
            "initial_rah_equals_ra_upper_descriptive": True,
            "initial_wjt_below_firm_lower_bound": True,
            "initial_household_wage_bound_status": "COMPOSITE_WAGE_HAS_NO_WJT_BOUND",
            "corrected_turn1_firm_ra0": f(corrected_row["firm_ra0"]),
            "corrected_turn1_firm_ra_used": f(corrected_row["firm_ra_used"]),
            "corrected_turn1_wage_raw": f(corrected_row["firm_wage_raw"]),
            "corrected_turn1_wage_used": f(corrected_row["firm_wage_used"]),
            "corrected_turn1_ra_clipped_lower": corrected_row["ra_clipped_lower"],
            "corrected_turn1_wage_clipped_lower": corrected_row["wage_clipped_lower"],
            "corrected_turn1_wage_clipped_upper": corrected_row["wage_clipped_upper"],
            "MATLAB_initial_K_over_Y_model_units": matlab_cap_transformed / matlab_gdp_transformed,
            "MATLAB_initial_K_over_L_model_units": matlab_cap_transformed / matlab_pop_transformed,
            "MATLAB_initial_Y_over_N_model_units": matlab_gdp_transformed / matlab_pop_transformed,
            "canonical_initial_K_over_Y_model_units": canonical_cap_transformed / canonical_gdp_transformed,
            "canonical_initial_K_over_L_model_units": canonical_cap_transformed / canonical_pop_transformed,
            "canonical_initial_Y_over_N_model_units": canonical_gdp_transformed / canonical_pop_transformed,
            "provenance": (
                "MATLAB cache mydata2{10}; GDP/CAP/POP row10=2009; IND_Zt fixed row21=2020; "
                "alpha PLM vintage19; corrected values canonical workbook 2018; saved outputs explicitly labelled by route"
            ),
        }
        ledger.append({key: finite_or_text(value) for key, value in row.items()})

    if len(ledger) != 31:
        raise AssertionError("ledger must contain 31 provinces")
    if not all(row["alpha_classification"] == "EXACT_MATCH" for row in ledger):
        raise AssertionError("expected alpha cache to match canonical vintage19")

    ledger_csv = output / "province_comparison_ledger.csv"
    ledger_md = output / "province_comparison_ledger.md"
    write_csv(ledger_csv, ledger)
    make_markdown_ledger(ledger_md, ledger)

    discrepancy_rows: list[dict[str, Any]] = []
    for metric in (
        "GDP_relative_difference",
        "POP_relative_difference",
        "capital_relative_difference",
        "Zt_relative_difference",
        "legacy_replay_turn1_GovInv_to_private_supply",
        "MATLAB_initial_K_over_Y_model_units",
        "MATLAB_initial_K_over_L_model_units",
        "MATLAB_initial_Y_over_N_model_units",
    ):
        discrepancy_rows.extend(top_rows(ledger, metric))
    write_csv(output / "largest_discrepancy_summary.csv", discrepancy_rows)

    lineage = [
        {"variable": "2018 label", "source_type": "runtime index", "source": "multi_prov_HANK_12sts.m:118-133", "operation": "year=ii+2008; data_MAT{ii}; data_year=ii", "unit": "index/year", "finding": "ii=10 labels 2018 but passes level row index 10"},
        {"variable": "GDP", "source_type": "processed workbook then cache", "source": "load_GDPdata.m:74,93,106-110; mpHANK_equilibrium_2000.m:31,41", "operation": "GDP workbook x1000; row data_year", "unit": "亿元 x1000 model units", "finding": "2018-labelled route consumes row10=2009"},
        {"variable": "CAP/Kt0", "source_type": "PIM workbook then cache", "source": "load_GDPdata.m:75,94; mpHANK_equilibrium_2000.m:27-28", "operation": "总资本存量 x1000; row data_year", "unit": "万元 x1000 model units", "finding": "2018-labelled route consumes row10=2009"},
        {"variable": "PIM capital construction", "source_type": "deterministic source transform", "source": "load_GDPdata.m:55-61; It_to_Kt.m:1-10", "operation": "K0=I0/.1; Kt=(1-delta)K(t-1)+I(t-1), delta=.096", "unit": "万元", "finding": "2009 capital uses investment through 2008"},
        {"variable": "POP/N", "source_type": "processed workbook then cache", "source": "load_GDPdata.m:79,104; mpHANK_equilibrium_2000.m:29-30", "operation": "常住人口 x100; row data_year", "unit": "万人 x100 model units", "finding": "2018-labelled route consumes row10=2009"},
        {"variable": "alpha", "source_type": "PLM workbook persisted in cache", "source": "load_GDPdata.m:115-118,131; mpHANK_equilibrium_2000.m:26", "operation": "last coefficient in 总面板回归系数_19_行业4", "unit": "elasticity", "finding": "vintage19 matches corrected canonical 2018"},
        {"variable": "Zt", "source_type": "cache derived from level data", "source": "load_GDPdata.m:126-137; mpHANK_equilibrium_2000.m:25", "operation": "GDP(row21)*CAP(row21)^(-alpha)*POP(row21)^(alpha-1)", "unit": "model productivity", "finding": "fixed row21=2020 for every mydata2 vintage; not same-year 2018"},
        {"variable": "GovInv", "source_type": "runtime initialization", "source": "multi_prov_HANK_12sts.m:83; mpHANK_equilibrium_2000.m:40", "operation": "Kt0*GovInv_ratio; ratio=1", "unit": "capital model units", "finding": "equals 2009 Kt0 in 2018-labelled route"},
        {"variable": "At/Bt", "source_type": "runtime initialization then household output", "source": "multi_prov_HANK_12sts.m:101-105; HANK_mp_1turn.m:15,31", "operation": "initial 2/1; household return required before allocation", "unit": "per household", "finding": "original MATLAB persisted 2018 st cache absent"},
        {"variable": "Kt_supply/rah", "source_type": "allocation output", "source": "HANK_mp_1turn.m:31-40", "operation": "uses household At*N and old ra vector", "unit": "capital / return", "finding": NOT_AVAILABLE},
        {"variable": "Kt_total/ra/wjt", "source_type": "firm output", "source": "HANK_firm.m:14,54-84", "operation": "Kt_supply+GovInv then firm equations and clips", "unit": "capital / rates", "finding": NOT_AVAILABLE},
        {"variable": "household wage", "source_type": "controller state", "source": "HANK_mp_1turn.m:50-53; wage_caculate.m:1-16", "operation": "destination wage aggregator after firm", "unit": "composite wage", "finding": "initial w=20; updated value unavailable without model rerun"},
        {"variable": "sigmau migration cost", "source_type": "distance workbook transform", "source": "load_distdata.m:1-8", "operation": "distance/max(distance)*max_sigmau", "unit": "migration-cost wedge", "finding": "static exogenous matrix; not a target of this numeric comparison"},
    ]
    write_csv(output / "data_source_lineage.csv", lineage)

    unit_audit = [
        {"object": "GDP", "raw_unit": "亿元", "MATLAB_transform": "x1000", "canonical_transform": "x1000", "assessment": "MATCH_AFTER_DOCUMENTED_UNIT_TRANSFORM; year differs 2009 vs 2018"},
        {"object": "POP", "raw_unit": "万人", "MATLAB_transform": "x100", "canonical_transform": "x100", "assessment": "MATCH_AFTER_DOCUMENTED_UNIT_TRANSFORM; year differs 2009 vs 2018"},
        {"object": "PIM capital", "raw_unit": "万元", "MATLAB_transform": "x1000", "canonical_transform": "x1000", "assessment": "same implemented scaling, but dimensional scale is not a documented 万元-to-亿元 conversion and year differs"},
        {"object": "Zt", "raw_unit": "derived", "MATLAB_transform": "fixed row21 with vintage-specific alpha", "canonical_transform": "same-year level row19 with vintage19 alpha", "assessment": "YEAR_OR_VINTAGE_MISMATCH"},
        {"object": "GovInv", "raw_unit": "capital model units", "MATLAB_transform": "Kt0*1", "canonical_transform": "Kt0*1", "assessment": "same rule; inherited capital year mismatch"},
        {"object": "initial wjt", "raw_unit": "firm wage state", "MATLAB_transform": "constant 0.6", "canonical_transform": "same source constant", "assessment": "below configured firm output lower bound 0.8 before any firm call"},
        {"object": "initial ra/rah", "raw_unit": "rate", "MATLAB_transform": "constant 0.09", "canonical_transform": "same source constant", "assessment": "ra at upper bound; rah only descriptively equal because rah has no native firm clip"},
    ]
    write_csv(output / "unit_scaling_audit.csv", unit_audit)

    role_map = [
        {"variable": "GDP/Y0", "role": "exogenous data target", "enters": "initial Yt/Yt0 and controller Zt target", "source_status": "2009 row in original 2018-labelled route"},
        {"variable": "POP/N", "role": "exogenous data", "enters": "household mass, At*N and migration labor", "source_status": "2009 row in original 2018-labelled route"},
        {"variable": "alpha", "role": "calibration/regression derived", "enters": "firm Cobb-Douglas and Zt construction", "source_status": "PLM vintage19"},
        {"variable": "Zt", "role": "calibration derived then controller-updated", "enters": "firm output/wage/rental rate", "source_status": "fixed 2020 level construction in original cache"},
        {"variable": "GovInv", "role": "initialized/controller-updated state", "enters": "Kt=Kt_supply+GovInv", "source_status": "initialized equal Kt0"},
        {"variable": "ra/rb/rah/w", "role": "initialized household price state", "enters": "household call", "source_status": "source constants; subsequent values require model calls"},
        {"variable": "At/Bt/Lt/Ct", "role": "household output", "enters": "allocation, migration, fiscal diagnostic", "source_status": "only accepted replay/corrected artifacts available"},
        {"variable": "Kt_supply", "role": "capital-allocation output", "enters": "firm capital", "source_status": NOT_AVAILABLE},
        {"variable": "firm ra/wjt", "role": "firm output with clipping", "enters": "next rah and composite wage", "source_status": NOT_AVAILABLE},
        {"variable": "next rah/w", "role": "allocation/controller state", "enters": "next household call", "source_status": NOT_AVAILABLE},
    ]
    write_csv(output / "outer_loop_variable_role_map.csv", role_map)

    source_paths = [
        source / name
        for name in (
            "main.m",
            "multi_prov_HANK_12sts.m",
            "mpHANK_equilibrium_2000.m",
            "load_GDPdata.m",
            "HANK_mp_1eq.m",
            "HANK_mp_1turn.m",
            "Lt_seperate.m",
            "HANK_firm.m",
            "wage_caculate.m",
            "It_to_Kt.m",
            "load_distdata.m",
        )
    ] + [raw, processed, plm, cache, distance_workbook, canonical]
    hash_receipt = {
        "schema": "CH5_MATLAB_2018_INPUT_SOURCE_HASH_RECEIPT_V1",
        "canonical_expected_sha256": EXPECTED_CANONICAL_SHA256,
        "canonical_actual_sha256": actual_canonical_hash,
        "canonical_match": True,
        "files": [
            {"path": str(path), "size": path.stat().st_size, "sha256": sha256(path)} for path in source_paths
        ],
        "missing_directly_referenced_cache": str(steady_cache),
        "missing_directly_referenced_cache_exists": steady_cache.exists(),
    }
    write_json(output / "source_data_hash_receipt.json", hash_receipt)

    call_ledger = {
        "schema": "CH5_MATLAB_2018_INPUT_AUDIT_ZERO_CALL_LEDGER_V1",
        "scientific_model_calls_total": 0,
        "matlab_model_calls": 0,
        "python_household_hjb_kfe_calls": 0,
        "firm_wage_migration_capital_allocation_calls": 0,
        "outer_turn_calls": 0,
        "steady_state_ge_annual_irf_results_calls": 0,
        "root_direct_iterative_eigen_scientific_solves": 0,
        "allowed_operations_only": [
            "static source inspection",
            "read-only XLSX/MAT/cache extraction",
            "hashing",
            "deterministic arithmetic",
            "CSV/JSON/Markdown serialization and static assertions",
        ],
    }
    write_json(output / "call_ledger.json", call_ledger)

    summary = {
        "schema": "CH5_MATLAB_2018_INPUT_AUDIT_SUMMARY_V1",
        "primary_verdict": "MATLAB_2018_INPUT_DATA_AUDIT_PASS__MATERIAL_EXTERNAL_SCALE_DIFFERENCES_IDENTIFIED",
        "province_count": 31,
        "matlab_2018_path_clear": True,
        "matlab_2018_actual_level_year": 2009,
        "matlab_2018_zt_level_year": 2020,
        "matlab_2018_plm_vintage": 19,
        "canonical_level_year": 2018,
        "canonical_plm_vintage": 19,
        "authority_document_inconsistency": "rule index/task active; status and handoff still say no active task",
        "steady_cache_2018_present": steady_cache.exists(),
        "cache_workbook_exact_checks": cache_workbook_diffs,
        "zt_row21_reconstruction_max_abs_diff": zt_reconstruction_max_abs_diff,
        "material_findings": [
            "2018 label mixes 2009 GDP/CAP/POP levels with fixed-2020 Zt and PLM-vintage19 alpha",
            "alpha matches canonical exactly, while GDP/POP/capital/Zt compare different years",
            "capital transform multiplies a source documented as 万元 by the same x1000 multiplier used for GDP亿元; no dimensional conversion rationale is documented in source",
            "GovInv starts equal to Kt0 and therefore inherits the capital year and scale",
            "initial ra is at its upper bound and initial wjt=0.6 is below the firm-output lower bound 0.8 before the first firm call",
        ],
        "unavailable_without_rerun": [
            "original MATLAB 2018 persisted At/Bt/Lt/Ct (steady cache absent)",
            "original MATLAB 2018 Kt_supply and Kt_total",
            "original MATLAB first-turn firm ra/wjt and next rah/household wage",
        ],
        "scientific_model_calls_total": 0,
        "causality_boundary": "External-input discrepancies are candidate propagation inputs only; this audit does not identify a cause of turn3 asset collapse.",
        "top_discrepancies": discrepancy_rows,
    }
    write_json(output / "audit_summary.json", summary)

    for name in (
        "audit_summary.json",
        "call_ledger.json",
        "data_source_lineage.csv",
        "largest_discrepancy_summary.csv",
        "outer_loop_variable_role_map.csv",
        "province_comparison_ledger.csv",
        "province_comparison_ledger.md",
        "source_data_hash_receipt.json",
        "unit_scaling_audit.csv",
    ):
        path = output / name
        target = evidence / name
        target.write_bytes(path.read_bytes())

    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", required=True)
    parser.add_argument("--matlab-root", required=True)
    parser.add_argument("--canonical", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--evidence", required=True)
    args = parser.parse_args()
    summary = build(args)
    print(json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
