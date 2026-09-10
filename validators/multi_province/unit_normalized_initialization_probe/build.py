from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
from pathlib import Path
from statistics import median

from openpyxl import load_workbook


VERDICT = "UNIT_NORMALIZED_INITIALIZATION_PROBE_PASS__DATA_CONSISTENT_PRICE_RECEIPT_COMPLETE"
ALPHA = 0.7380939146868483
ALPHA_BOUNDS = (0.2, 0.8)
RA_BOUNDS = (0.02, 0.09)
WAGE_BOUNDS = (0.8, 1.3)
PARAM = {
    "epsilon": 10.0, "theta": 100.0, "delta_firm": 0.025,
    "pit": 0.02, "pit_1": 0.02, "rk_old": 0.1, "corptau": 0.25,
    "tau": 0.05, "max_phi": 0.3, "max_sigmau": 0.5,
    "phi_l": 5.0, "alphal": 1.0, "legacy_At": 2.0,
}
SOURCE_HASHES = {
    "HANK_firm.m": "EE02C15414ADF9F99AADE04F1F22E64FA7094C8AB77753B6130BC4BFA6CE7BD5",
    "HANK_mp_1turn.m": "D3D03F37286ED66202673EA63D49BABCE8D5309BAC9C13793C8E60585C21FECF",
    "wage_caculate.m": "0FB84B51E2BE50CD3D065D33385882311A31E12596AEEB0CE2C808A8C36B6A63",
    "mpHANK_equilibrium_2000.m": "26EA44552DA33919F8CCD777C084E15ECA0EA9575FEE80A07F9E0056F3F97DE5",
    "load_GDPdata.m": "DECA8AF3F22097550B8957FE848989E6342619CB9929A1C00076E020549366C5",
    "load_distdata.m": "18F594DD7D1ED090CA2AF576DEBCD8DCAA73C012608A8921F8D5BD6CC24F478B",
    "multi_prov_HANK_12sts.m": "3C44449CFD4047B5C9E17E540AFEA2F50B4251150F8F74AB8CCEED26E15DEC97",
    "中国各省省会地理距离矩阵.xlsx": "26E44D174A8EFFBDCA526D95DA38F0E5883E0C78FDFD036D2DFF1D1FBA5A3566",
}
PROVINCE_ORDER = [
    "北京", "天津", "河北", "山西", "内蒙古", "辽宁", "吉林", "黑龙江", "上海", "江苏",
    "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南", "广东", "广西",
    "海南", "重庆", "四川", "贵州", "云南", "西藏", "陕西", "甘肃", "青海", "宁夏", "新疆",
]
DISTANCE_LABEL_ORDER = [
    "北京市", "天津市", "河北省", "山西省", "内蒙古", "辽宁省", "吉林省", "黑龙江", "上海市", "江苏省",
    "浙江省", "安徽省", "福建省", "江西省", "山东省", "河南省", "湖北省", "湖南省", "广东省", "广西",
    "海南省", "重庆市", "四川省", "贵州省", "云南省", "西藏", "陕西省", "甘肃省", "青海省", "宁夏", "新疆",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def classify(value: float, bounds: tuple[float, float]) -> str:
    return "below" if value < bounds[0] else "above" if value > bounds[1] else "inside"


def clipped(value: float, bounds: tuple[float, float]) -> float:
    return min(max(value, bounds[0]), bounds[1])


def load_distance(path: Path) -> list[list[float]]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["geom"]
    column_labels = [sheet.cell(1, col).value for col in range(2, 33)]
    row_labels = [sheet.cell(row, 1).value for row in range(2, 33)]
    if column_labels != DISTANCE_LABEL_ORDER or row_labels != DISTANCE_LABEL_ORDER:
        raise ValueError("distance workbook province order does not match the frozen 31-province source order")
    values = [[sheet.cell(row, col).value for col in range(2, 33)] for row in range(2, 33)]
    book.close()
    if len(values) != 31 or any(len(row) != 31 for row in values):
        raise ValueError("distance workbook geom numeric body is not 31x31")
    matrix = [[float(value) for value in row] for row in values]
    if any(not math.isfinite(value) or value < 0 for row in matrix for value in row):
        raise ValueError("distance matrix contains invalid values")
    if any(abs(matrix[i][i]) > 1e-12 for i in range(31)):
        raise ValueError("distance matrix diagonal is not zero")
    if any(abs(matrix[i][j] - matrix[j][i]) > 1e-9 for i in range(31) for j in range(31)):
        raise ValueError("distance matrix is not symmetric")
    maximum = max(max(row) for row in matrix)
    if maximum <= 0:
        raise ValueError("distance maximum must be positive")
    return [[value / maximum * PARAM["max_sigmau"] for value in row] for row in matrix]


def stats(values: list[float]) -> dict[str, float]:
    return {"min": min(values), "median": median(values), "max": max(values)}


def markdown_table(rows: list[dict], fields: list[str]) -> str:
    lines = ["| " + " | ".join(fields) + " |", "| " + " | ".join(["---"] * len(fields)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(row[field]) for field in fields) + " |")
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", type=Path, required=True)
    parser.add_argument("--protected-root", type=Path, required=True)
    parser.add_argument("--evidence-root", type=Path, required=True)
    parser.add_argument("--session-build-processes", type=int, default=1)
    parser.add_argument("--session-focused-test-processes", type=int, default=1)
    parser.add_argument("--session-compile-processes", type=int, default=0)
    parser.add_argument("--session-finalize-processes", type=int, default=1)
    args = parser.parse_args()
    repo = args.repo_root.resolve()
    protected = args.protected_root.resolve()
    evidence = args.evidence_root.resolve()
    if evidence.exists() and any(evidence.iterdir()):
        raise FileExistsError(f"refusing to overwrite non-empty evidence root: {evidence}")
    evidence.mkdir(parents=True, exist_ok=True)

    source_receipts = []
    for name, expected in SOURCE_HASHES.items():
        path = protected / name
        actual = sha256(path)
        if actual != expected:
            raise ValueError(f"protected source hash mismatch: {name}: {actual}")
        source_receipts.append({"file": name, "path": str(path), "sha256": actual, "match": True})

    ledger_path = repo / "reports/mp4c_2018_raw_nbs_rebuild_20260910/corrected_2018_vs_matlab_ledger.csv"
    with ledger_path.open("r", encoding="utf-8-sig", newline="") as handle:
        input_rows = list(csv.DictReader(handle))
    if (len(input_rows) != 31 or [int(row["province_index"]) for row in input_rows] != list(range(1, 32))
            or [row["province_name"] for row in input_rows] != PROVINCE_ORDER):
        raise ValueError("accepted ledger does not contain the ordered 31-province contract")
    if any(not math.isclose(float(row["new_alpha"]), ALPHA, rel_tol=0, abs_tol=1e-14) for row in input_rows):
        raise ValueError("accepted alpha is not uniform or does not match the frozen task value")

    mstar = 1.0 - 1.0 / PARAM["epsilon"]
    mt0 = PARAM["rk_old"] * PARAM["pit"] * PARAM["theta"] / PARAM["epsilon"] + mstar
    rows: list[dict] = []
    for src in input_rows:
        y_raw = float(src["corrected_raw_nbs_gdp_2018_100m_yuan"])
        n_raw = float(src["corrected_raw_nbs_population_2018_10k_persons"])
        k_raw = float(src["corrected_track_a_capital_2018_100m_yuan"])
        y, n, k = y_raw * 1_000.0, n_raw * 100.0, k_raw * 1_000.0
        if min(y, n, k) <= 0:
            raise ValueError(f"non-positive input for {src['province_name']}")
        z = y / (k**ALPHA * n ** (1.0 - ALPHA))
        raw_wage = mt0 * (1.0 - ALPHA) * z * (k / n) ** ALPHA
        rk_new = mt0 * ALPHA * y / k
        theta_cost = PARAM["theta"] / 2.0 * PARAM["pit"] ** 2 * y
        profit = max((1.0 - mt0) * y - theta_cost, 0.0)
        corptax_before = profit * PARAM["corptau"]
        dividend_rate = profit * (1.0 - PARAM["corptau"]) / k
        raw_ra = rk_new - PARAM["delta_firm"] + dividend_rate
        ra_status, wage_status = classify(raw_ra, RA_BOUNDS), classify(raw_wage, WAGE_BOUNDS)
        used_ra, used_wage = clipped(raw_ra, RA_BOUNDS), clipped(raw_wage, WAGE_BOUNDS)
        ra_tax_adjustment = (raw_ra - used_ra) * k
        wage_tax_adjustment = (raw_wage - used_wage) * n
        corptax_after_ra = corptax_before + ra_tax_adjustment
        corptax_after_wage = corptax_after_ra + wage_tax_adjustment
        rows.append({
            "province_index": int(src["province_index"]), "province_name": src["province_name"],
            "Y0_raw_亿元": y_raw, "Y0_MU": y, "POP_raw_万人": n_raw, "N0_NU": n,
            "K0_trackA_亿元": k_raw, "K0_MU": k, "L0": n,
            "L0_proxy_role": "2018_POPULATION_AS_LABOR_PROXY_NU",
            "alpha_raw": ALPHA, "alpha_used": ALPHA, "alpha_clip_flag": False,
            "alpha_clip_reason": "NONE_INSIDE_OWNER_RANGE_0.2_0.8", "Zt0": z,
            "Kt_1": k, "Lt_1": n, "Yt_1": y, "Zt_1": z,
            "pit": PARAM["pit"], "pit_1": PARAM["pit_1"], "rk_old": PARAM["rk_old"],
            "mstar": mstar, "mt0": mt0,
            "raw_wjt0": raw_wage, "used_wjt0": used_wage,
            "wage_bound_status": wage_status, "wage_lower_clip": wage_status == "below",
            "wage_upper_clip": wage_status == "above",
            "rk_new": rk_new, "Thetat": theta_cost, "PIt": profit,
            "Corptax_before_clipping": corptax_before, "divrate": dividend_rate,
            "raw_ra0": raw_ra, "used_ra0": used_ra, "ra_bound_status": ra_status,
            "ra_lower_clip": ra_status == "below", "ra_upper_clip": ra_status == "above",
            "Corptax_ra_adjustment": ra_tax_adjustment,
            "Corptax_after_ra_clip": corptax_after_ra,
            "Corptax_wage_adjustment": wage_tax_adjustment,
            "Corptax_after_ra_and_wage_clip": corptax_after_wage,
            "legacy_ra": 0.09, "legacy_wjt": 0.6,
            "raw_ra_minus_legacy": raw_ra - 0.09, "raw_wjt_minus_legacy": raw_wage - 0.6,
            "legacy_matlab_level_year": int(src["legacy_matlab_level_year"]),
            "legacy_matlab_gdp_亿元": float(src["legacy_matlab_gdp_100m_yuan"]),
            "legacy_matlab_population_万人": float(src["legacy_matlab_population_10k_persons"]),
            "legacy_matlab_capital_亿元": float(src["legacy_matlab_capital_100m_yuan"]),
            "K_to_Y": k / y, "K_to_N": k / n, "Y_to_N": y / n,
            "firm_equation_provenance": "HANK_firm.m:30,33-35,43-54",
            "clipping_provenance": "HANK_firm.m:55-74",
            "unit_provenance": "exact task section 3: GDP/K x1000 MU; POP x100 NU",
        })

    pcap = [row["K_to_N"] for row in rows]
    low, high = min(pcap), max(pcap)
    ratios = [0.3 * (value - low) / (high - low) for value in pcap]
    tempra = sum(ratio * row["used_ra0"] for ratio, row in zip(ratios, rows))
    sigmau = load_distance(protected / "中国各省省会地理距离矩阵.xlsx")
    productivity = [row["Y_to_N"] for row in rows]
    for i, row in enumerate(rows):
        ratio = ratios[i]
        row["inter_prv_ratio"] = ratio
        row["rah0"] = ((1.0 - ratio) * row["used_ra0"]
                       + ratio / 30.0 * (tempra - ratio * row["used_ra0"]))
        row["rah0_provenance"] = "HANK_mp_1turn.m:29-40; SOURCE_LAGGED_BASELINE_INITIAL_SEED"
        total = 0.0
        for j in range(31):
            phi = 1.0 + PARAM["max_phi"] * (productivity[j] - productivity[i]) / (productivity[j] + productivity[i])
            base = rows[j]["used_wjt0"] * (1.0 - PARAM["tau"] - sigmau[j][i]) / phi
            if base <= 0:
                raise ValueError("static wage base is non-positive")
            total += phi * base ** (1.0 + 1.0 / PARAM["phi_l"])
        row["w0"] = PARAM["alphal"] ** (-PARAM["phi_l"] / (1.0 + PARAM["phi_l"])) * total ** (PARAM["phi_l"] / (1.0 + PARAM["phi_l"]))
        row["w0_provenance"] = "HANK_mp_1turn.m:4-10; wage_caculate.m:3-11; load_distdata.m:6-8"
        row["composite_seed_policy"] = "INITIAL_COMPOSITE_NO_DAMPING_BASE"

    out = repo / "reports/mp4c_unit_normalized_initialization_probe_20260910"
    out.mkdir(parents=True, exist_ok=True)
    receipt_path = out / "province_initialization_receipt.csv"
    with receipt_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
    selected = ["province_index", "province_name", "Y0_MU", "N0_NU", "K0_MU", "Zt0", "raw_wjt0", "used_wjt0", "wage_bound_status", "raw_ra0", "used_ra0", "ra_bound_status", "rah0", "w0", "K_to_Y", "K_to_N"]
    (out / "province_initialization_receipt.md").write_text(
        "# Province initialization receipt\n\n" + markdown_table(rows, selected), encoding="utf-8")

    bridge = []
    for row in rows:
        atn = PARAM["legacy_At"] * row["N0_NU"]
        bridge.append({
            "province_index": row["province_index"], "province_name": row["province_name"],
            "legacy_At_grid": PARAM["legacy_At"], "N0_NU": row["N0_NU"],
            "AtN_baseline_MU": atn, "K0_MU": row["K0_MU"],
            "GovInv_source_initial_reference_MU": row["K0_MU"],
            "AtN_to_K0": atn / row["K0_MU"], "AtN_to_GovInv_reference": atn / row["K0_MU"],
            "bridge_status": "SOURCE_FAITHFUL_BASELINE_ONLY",
            "authority": "NO_REPLACEMENT_BRIDGE_SELECTED",
        })
    with (out / "asset_bridge_scale_diagnostic.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(bridge[0])); writer.writeheader(); writer.writerows(bridge)

    def counts(field: str, bounds: tuple[float, float]) -> dict[str, int]:
        values = [row[field] for row in rows]
        return {key: sum(classify(value, bounds) == key for value in values) for key in ("below", "inside", "above")}
    raw_ra = [row["raw_ra0"] for row in rows]; used_ra = [row["used_ra0"] for row in rows]
    raw_w = [row["raw_wjt0"] for row in rows]; used_w = [row["used_wjt0"] for row in rows]
    any_clip = sum(row["ra_bound_status"] != "inside" or row["wage_bound_status"] != "inside" for row in rows)
    summary = {
        "schema": "CH5_UNIT_NORMALIZED_INITIALIZATION_SUMMARY_V1", "verdict": VERDICT,
        "province_count": 31, "year": 2018, "capital_route": "RAW_NBS_GFCF_TRACK_A_PIM_DELTA_0.096",
        "units": {"MU": "10万元", "NU": "100 persons", "GDP_亿元_multiplier": 1000, "capital_亿元_multiplier": 1000, "POP_万人_multiplier": 100},
        "mt0": mt0, "mt0_note": "Protected HANK_firm.m:35 retains rk_old*pit*theta/epsilon; this is .92, not the accepted redesign prose simplification .9.",
        "raw_ra_bound_counts": counts("raw_ra0", RA_BOUNDS), "raw_wage_bound_counts": counts("raw_wjt0", WAGE_BOUNDS),
        "raw_ra_stats": stats(raw_ra), "used_ra_stats": stats(used_ra),
        "raw_wage_stats": stats(raw_w), "used_wage_stats": stats(used_w),
        "rah0_stats": stats([row["rah0"] for row in rows]), "w0_stats": stats([row["w0"] for row in rows]),
        "any_firm_price_clip_provinces": any_clip,
        "mostly_on_clipping_boundaries": any_clip > 31 / 2,
        "largest_raw_ra_discrepancy_from_legacy": max(rows, key=lambda row: abs(row["raw_ra_minus_legacy"])),
        "largest_raw_wage_discrepancy_from_legacy": max(rows, key=lambda row: abs(row["raw_wjt_minus_legacy"])),
        "highest_K_to_Y": max(rows, key=lambda row: row["K_to_Y"]), "lowest_K_to_Y": min(rows, key=lambda row: row["K_to_Y"]),
        "highest_K_to_N": max(rows, key=lambda row: row["K_to_N"]), "lowest_K_to_N": min(rows, key=lambda row: row["K_to_N"]),
        "anhui": next(row for row in rows if row["province_name"] == "安徽"),
        "results_eligible": False,
    }
    write_json(out / "national_initialization_summary.json", summary)
    equations = {
        "schema": "CH5_UNIT_NORMALIZED_SOURCE_EQUATION_RECEIPT_V1",
        "accepted_input_ledger": {"path": str(ledger_path), "sha256": sha256(ledger_path), "rows": 31},
        "protected_sources": source_receipts,
        "equations": {
            "Zt0": "Y0/(K0^alpha_used*L0^(1-alpha_used))",
            "mt0": "rk_old*pit*theta/epsilon + (1-1/epsilon) = 0.92 under zero changes and pit=pit_1",
            "raw_wjt0": "mt0*(1-alpha)*Zt0*(K0/L0)^alpha",
            "raw_ra0": "mt0*alpha*Y0/K0-delta_firm+max((1-mt0)*Y0-Thetat,0)*(1-corptau)/K0",
            "Thetat": "theta/2*pit^2*Y0",
            "Corptax_ra_adjustment": "(raw_ra0-used_ra0)*K0",
            "Corptax_wage_adjustment": "(raw_wjt0-used_wjt0)*L0",
            "inter_prv_ratio": "0.3*(K0/N0-min(K0/N0))/(max(K0/N0)-min(K0/N0))",
            "rah0": "(1-ratio_i)*used_ra_i + ratio_i/30*(sum_j ratio_j*used_ra_j-ratio_i*used_ra_i)",
            "phi(j,i)": "1+0.3*((Yj/Lj)-(Yi/Li))/((Yj/Lj)+(Yi/Li))",
            "sigmau(j,i)": "distance(j,i)/max(distance)*0.5",
            "w0": "alphal^(-phi_l/(1+phi_l))*sum_j[phi(j,i)*(used_wjt_j*(1-tau-sigmau(j,i))/phi(j,i))^(1+1/phi_l)]^(phi_l/(1+phi_l))",
        },
        "spec_discrepancy": {"accepted_redesign_prose_mt0": 0.9, "protected_source_mt0": mt0, "resolution": "SOURCE_EQUATION_GOVERNS_THIS_PROBE__NO_SOURCE_MODIFICATION"},
    }
    write_json(out / "source_equation_receipt.json", equations)
    calls = {
        "schema": "CH5_UNIT_NORMALIZED_INITIALIZATION_CALL_LEDGER_V1",
        "accounting_scope": "CUMULATIVE_CURRENT_TASK_SESSION_INCLUDING_PRESERVED_PRELIMINARY_EVIDENCE_RUNS",
        "deterministic_helper_processes": args.session_build_processes,
        "firm_equation_row_evaluations": 31 * args.session_build_processes,
        "return_composite_row_evaluations": 31 * args.session_build_processes,
        "wage_composite_row_evaluations": 31 * args.session_build_processes,
        "focused_test_processes": args.session_focused_test_processes,
        "focused_test_cases_executed": 6 * args.session_focused_test_processes,
        "python_compile_processes": args.session_compile_processes,
        "evidence_finalize_processes": args.session_finalize_processes,
        "preserved_preliminary_evidence_roots": [
            "D:/ProjectTemp/ch5-unit-normalized-initialization-probe-20260910-001",
            "D:/ProjectTemp/ch5-unit-normalized-initialization-probe-20260910-002",
            "D:/ProjectTemp/ch5-unit-normalized-initialization-probe-20260910-003",
        ],
        "matlab": 0, "household_hjb": 0, "kfe": 0, "household_control_solve": 0,
        "labor_allocation_Lt_seperate": 0, "outer_turn": 0, "steady_state": 0,
        "root_newton_broyden_fsolve_brent_anderson": 0, "ge": 0, "annual": 0,
        "irf": 0, "results": 0, "parameter_tuning": 0,
    }
    write_json(out / "call_ledger.json", calls)
    write_json(evidence / "run_receipt.json", {"verdict": VERDICT, "output_directory": str(out), "call_ledger": calls})
    shutil.copy2(receipt_path, evidence / receipt_path.name)
    shutil.copy2(out / "national_initialization_summary.json", evidence / "national_initialization_summary.json")
    print(json.dumps({"verdict": VERDICT, "output": str(out), "evidence": str(evidence), "mt0": mt0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
